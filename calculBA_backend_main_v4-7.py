"""calculBA — version Streamlit (un seul fichier).

Application de collecte et d'analyse de données de génie civil : jeux de
données, saisie et import CSV/Excel, statistiques, graphiques, lois de
probabilité, exports.

ACCÈS : sans compte. Chaque visiteur reçoit un espace privé, lié à son lien
personnel (paramètre « ?espace= » de l'adresse) : il ne voit que son travail.

SUR VOTRE PC
    pip install -r requirements.txt
    streamlit run streamlit_app.py
    (ouvre http://localhost:8501 ; les données vont dans calculba.db, à côté de ce fichier)

SUR STREAMLIT COMMUNITY CLOUD (lien public gratuit)
    1. Dépôt GitHub contenant ce fichier et requirements.txt.
    2. share.streamlit.io : Main file path = streamlit_app.py
    3. Paramètres de l'application, « Secrets » : collez
           DATABASE_URL = "postgresql://...neon.tech/neondb?sslmode=require"
       Sans DATABASE_URL, les données sont EFFACÉES à chaque redémarrage de
       l'application : la base Neon (gratuite) les conserve.

Sommaire (cherchez « SECTION » avec Ctrl + F) :
    1. RÉGLAGES
    2. CALCULS (lecture des fichiers, statistiques, graphiques, lois, exports)
    3. BASE DE DONNÉES (SQLite sur PC, PostgreSQL en ligne)
    4. ESPACES PERSONNELS (SANS COMPTE)
    5. OPÉRATIONS SUR LES DONNÉES
    6. INTERFACE STREAMLIT
"""
from __future__ import annotations

import csv
import hashlib
import html
import io
import json
import math
import os
import re
import secrets
import sqlite3
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from datetime import time as dtime  # heure seule (cellules Excel) ; « time » désigne le module
from pathlib import Path
from typing import Any, Iterable, Sequence

import altair as alt
import numpy as np
import pandas as pd
import streamlit as st

# ============================================================================
# SECTION 1 — RÉGLAGES
# ============================================================================


def _setting(name: str, default: str = "") -> str:
    """Réglage lu dans les variables d'environnement, puis dans les « Secrets » Streamlit."""
    value = os.environ.get(name, "").strip()
    if value:
        return value
    try:
        return str(st.secrets.get(name, default)).strip()
    except Exception:  # pas de fichier de secrets (PC)
        return default


APP_NAME = "calculBA"
VERSION = "6.0.0-streamlit"
ROOT = Path(__file__).resolve().parent
DB_PATH = Path(_setting("CALCULBA_DB", str(ROOT / "calculba.db")))
DATABASE_URL = _setting("DATABASE_URL")
USE_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))

# Espaces personnels : un visiteur = un espace, identifié par un jeton dans l'adresse.
SPACE_PARAM = "espace"
SPACE_DOMAIN = "@espace.calculba.local"
SPACE_IDLE_DAYS = int(_setting("CALCULBA_SPACE_IDLE_DAYS", "180"))   # espace inutilisé : supprimé
PUBLIC_URL = _setting("CALCULBA_PUBLIC_URL", "https://calculba-w7t3fmcgaokaqxftwds8ay.streamlit.app/")

MAX_COLUMNS = 60
MAX_ROWS = 50_000
MAX_UPLOAD_BYTES = 15 * 1024 * 1024
MAX_NAME_LENGTH = 120
MAX_COLUMN_NAME_LENGTH = 80
MAX_VALUE_LENGTH = 500
MAX_DATASETS_PER_USER = int(_setting("CALCULBA_MAX_DATASETS", "60"))
MAX_RECORDS_PER_USER = int(_setting("CALCULBA_MAX_RECORDS", "100000"))
# Limite pour toute la base (protège la base gratuite Neon de 0,5 Go).
MAX_RECORDS_TOTAL = int(_setting("CALCULBA_MAX_RECORDS_TOTAL", "600000"))
NUMERIC_THRESHOLD = 0.8

DOMAINS = ["Béton", "Sols et géotechnique", "Topographie", "Matériaux", "Structures", "Routes", "Hydraulique", "Autre"]
EXAMPLES = {
    "Béton": "Echantillon, Age_jours, Resistance_MPa, Affaissement_mm",
    "Sols et géotechnique": "Sondage, Profondeur_m, Teneur_eau_pct, Densite_seche_t_m3, CBR_pct",
    "Topographie": "Point, Est_m, Nord_m, Altitude_m",
    "Matériaux": "Lot, Tamis_mm, Passant_pct",
    "Structures": "Element, Portee_m, Charge_kN, Fleche_mm",
    "Routes": "PR, Couche, Epaisseur_cm, Deflexion_1_100mm",
    "Hydraulique": "Station, Date, Debit_m3_s, Hauteur_m",
}
DATA_COLOR = "#2F63AE"


# ============================================================================
# SECTION 2 — CALCULS (repris tels quels de la version testée de calculBA)
# ============================================================================

_SPACES = re.compile(r"[\s  ']")  # espaces, insécables, apostrophe suisse


_NUMBER = re.compile(r"[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?")


# Caractères de contrôle (ex. Ctrl-Z de fin de fichier DOS) : invisibles, et un
# classeur .xlsx qui en contient est refusé par openpyxl. Tab, saut de ligne et
# retour chariot sont conservés.
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def to_number(value: Any) -> float | None:
    """Convertit une valeur saisie en nombre, ou None si ce n'en est pas un.

    Gère la virgule décimale (« 22,5 »), les séparateurs de milliers
    (« 1 234,5 », « 1.234,5 », « 1,234.5 ») et la notation scientifique.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, np.integer, np.floating)):
        f = float(value)
        return f if math.isfinite(f) else None
    s = _SPACES.sub("", str(value))
    if not s:
        return None
    if "," in s and "." in s:
        # Le séparateur le plus à droite est le séparateur décimal.
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(",") == 1:
        s = s.replace(",", ".")
    if not _NUMBER.fullmatch(s):
        return None
    f = float(s)
    return f if math.isfinite(f) else None


def is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not _CTRL.sub("", value).strip())


def clean_value(value: Any) -> Any:
    """Normalise une valeur avant stockage (sans la convertir)."""
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        return "VRAI" if value else "FAUX"
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        f = float(value)
        if not math.isfinite(f):
            return None
        return int(f) if f.is_integer() and abs(f) < 1e15 else f
    if isinstance(value, int):
        return value
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat(sep=" ") if (value.hour or value.minute or value.second) else value.date().isoformat()
    if isinstance(value, (date, dtime)):
        return value.isoformat()
    s = _CTRL.sub("", str(value)).strip()
    if not s:
        return None
    return s[: MAX_VALUE_LENGTH]


class ImportProblem(ValueError):
    """Erreur lisible renvoyée telle quelle à l'utilisateur."""


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _sniff_delimiter(sample: str) -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
    except csv.Error:
        # Repli : le séparateur le plus fréquent sur la ligne d'en-tête.
        first = sample.splitlines()[0] if sample else ""
        return max(";,\t|", key=first.count)


def read_table(filename: str, raw: bytes) -> tuple[list[str], list[list[Any]]]:
    """Retourne (colonnes, lignes) à partir d'un CSV ou d'un classeur .xlsx."""
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        text = _decode(raw)
        if not text.strip():
            raise ImportProblem("Le fichier est vide.")
        delimiter = _sniff_delimiter(text[:20_000])
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error as exc:  # cellule démesurée (> 131 072 caractères), etc.
            raise ImportProblem(
                "Ce fichier CSV ne peut pas être lu (cellule démesurée ou caractère invalide). "
                "Vérifiez-le, ou enregistrez-le à nouveau en .csv ou .xlsx."
            ) from exc
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            frame = pd.read_excel(io.BytesIO(raw), header=None, dtype=object, engine="openpyxl")
        except Exception as exc:  # fichier corrompu, protégé par mot de passe…
            raise ImportProblem("Ce classeur Excel ne peut pas être lu. Enregistrez-le à nouveau au format .xlsx.") from exc
        frame = frame.astype(object).where(pd.notna(frame), None)
        rows = frame.values.tolist()
    elif name.endswith(".xls"):
        raise ImportProblem("L'ancien format .xls n'est pas pris en charge. Dans Excel : Fichier › Enregistrer sous › Classeur Excel (.xlsx).")
    else:
        raise ImportProblem("Format non reconnu. Importez un fichier .csv ou .xlsx.")

    # Supprime les lignes entièrement vides (fréquentes en fin de feuille).
    rows = [r for r in rows if any(not is_empty(v) for v in r)]
    if not rows:
        raise ImportProblem("Le fichier ne contient aucune donnée.")

    header = [clean_value(v) for v in rows[0]]
    body = rows[1:]
    width = max(len(header), *(len(r) for r in body)) if body else len(header)
    header += [None] * (width - len(header))

    # Colonnes sans en-tête : acceptées seulement si elles sont entièrement vides.
    keep: list[int] = []
    for i, h in enumerate(header):
        column_empty = all(i >= len(r) or is_empty(r[i]) for r in body)
        if h is None:
            if column_empty:
                continue
            raise ImportProblem(f"La colonne n° {i + 1} contient des valeurs mais n'a pas de nom sur la première ligne.")
        keep.append(i)

    columns = [str(header[i]) for i in keep]
    validate_columns(columns)
    if not body:
        raise ImportProblem("Le fichier ne contient que la ligne d'en-tête.")
    if len(body) > MAX_ROWS:
        raise ImportProblem(f"Le fichier contient {len(body):,} lignes ; la limite est de {MAX_ROWS:,}.".replace(",", " "))

    data = [[clean_value(r[i]) if i < len(r) else None for i in keep] for r in body]
    return columns, data


def validate_columns(columns: list[str]) -> list[str]:
    cleaned = [str(c).strip() for c in columns]
    if not cleaned:
        raise ImportProblem("Ajoutez au moins une colonne.")
    if len(cleaned) > MAX_COLUMNS:
        raise ImportProblem(f"{len(cleaned)} colonnes trouvées ; la limite est de {MAX_COLUMNS}.")
    seen: dict[str, str] = {}
    for c in cleaned:
        if not c:
            raise ImportProblem("Chaque colonne doit avoir un nom.")
        if len(c) > MAX_COLUMN_NAME_LENGTH:
            raise ImportProblem(f"Le nom de colonne « {c[:30]}… » dépasse {MAX_COLUMN_NAME_LENGTH} caractères.")
        key = c.casefold()
        if key in seen:
            raise ImportProblem(f"La colonne « {c} » apparaît deux fois. Chaque nom de colonne doit être unique.")
        seen[key] = c
    return cleaned


def _r(x: float | None, nd: int = 4, sig: bool = True) -> float | None:
    if x is None or not math.isfinite(x):
        return None
    # Petites valeurs (perméabilité 1e-8 m/s, déformations…) : arrondir à nd décimales
    # les écraserait à 0. En dessous de ce seuil on garde 6 chiffres significatifs.
    # Les valeurs courantes (>= 0,01 pour nd = 4) sont arrondies comme avant.
    if sig and nd >= 4 and x != 0 and abs(x) < 10 ** (2 - nd):
        return float(f"{x:.6g}")
    return float(round(x, nd))


def _is_padded_code(filled: list) -> bool:
    """Codes à largeur fixe avec zéros de tête (001, 002… 120) : des identifiants, pas des mesures."""
    strs = [str(v).strip() for v in filled]
    if not all(re.fullmatch(r"[0-9]+", t) for t in strs):
        return False
    return any(len(t) > 1 and t.startswith("0") for t in strs) and len({len(t) for t in strs}) == 1


def profile(columns: list[str], rows: list[dict]) -> dict:
    """Type de chaque colonne et valeurs numériques extraites (alignées sur les lignes)."""
    out = {}
    for col in columns:
        raw = [r.get(col) for r in rows]
        filled = [v for v in raw if not is_empty(v)]
        nums = [to_number(v) for v in raw]
        parsed = sum(1 for v in nums if v is not None)
        numeric = (bool(filled) and parsed / len(filled) >= NUMERIC_THRESHOLD
                   and not _is_padded_code(filled))
        out[col] = {"numeric": numeric, "raw": raw, "nums": nums, "filled": len(filled), "parsed": parsed}
    return out


def histogram(values: np.ndarray) -> dict:
    if values.size == 0:
        return {"edges": [], "counts": []}
    lo, hi = float(values.min()), float(values.max())
    if lo == hi:
        return {"edges": [_r(lo - 0.5), _r(hi + 0.5)], "counts": [int(values.size)]}
    edges = np.histogram_bin_edges(values, bins="auto")
    n = int(min(30, max(5, len(edges) - 1)))
    counts, edges = np.histogram(values, bins=n)
    return {"edges": [_r(e, 6) for e in edges], "counts": [int(c) for c in counts]}


def describe(columns: list[str], rows: list[dict]) -> dict:
    prof = profile(columns, rows)
    total = len(rows)
    result: dict[str, dict] = {}
    for col in columns:
        p = prof[col]
        missing = total - p["filled"]
        if p["numeric"]:
            v = np.array([x for x in p["nums"] if x is not None], dtype=float)
            n = int(v.size)
            mean = float(v.mean()) if n else None
            std = float(v.std(ddof=1)) if n > 1 else None
            q1, med, q3 = (np.percentile(v, [25, 50, 75]).tolist() if n else (None, None, None))
            cv = (std / abs(mean) * 100) if (std is not None and mean) else None
            result[col] = {
                "type": "numeric",
                "count": n,
                "distinct": int(np.unique(v).size),
                "missing": missing,
                "invalid": p["filled"] - p["parsed"],
                # Numéros de ligne (à partir de 1, ordre de saisie) des valeurs non numériques,
                # pour que l'interface puisse y mener directement.
                "invalid_rows": [i + 1 for i, (raw, num) in enumerate(zip(p["raw"], p["nums"]))
                                 if num is None and not is_empty(raw)][:20],
                "mean": _r(mean), "std": _r(std), "cv": _r(cv, 2),
                "min": _r(float(v.min())) if n else None,
                "q1": _r(q1), "median": _r(med), "q3": _r(q3),
                "max": _r(float(v.max())) if n else None,
                "histogram": histogram(v),
            }
        else:
            values = [str(x).strip() for x in p["raw"] if not is_empty(x)]
            counts = pd.Series(values, dtype="string").value_counts() if values else pd.Series(dtype="int64")
            result[col] = {
                "type": "text",
                "count": len(values),
                "missing": missing,
                "distinct": int(counts.size),
                "top": [{"value": str(k), "count": int(c)} for k, c in counts.head(5).items()],
            }
    return {
        "total_rows": total,
        "numeric_columns": [c for c in columns if prof[c]["numeric"]],
        "text_columns": [c for c in columns if not prof[c]["numeric"]],
        "columns": result,
    }


def _thin(indices: list[int], limit: int) -> list[int]:
    """Échantillonnage régulier (déterministe) pour limiter le nombre de points tracés."""
    if len(indices) <= limit:
        return indices
    step = len(indices) / limit
    return [indices[int(i * step)] for i in range(limit)]


def series(rows: list[dict], column: str, limit: int = 1500) -> dict:
    nums = [to_number(r.get(column)) for r in rows]
    present = [i for i, v in enumerate(nums) if v is not None]
    kept = _thin(present, limit)
    return {
        "column": column,
        "points": [{"x": i + 1, "y": _r(nums[i], 6)} for i in kept],
        "plotted": len(kept),
        "available": len(present),
        "total_rows": len(rows),
    }


def scatter(rows: list[dict], x: str, y: str, limit: int = 2000) -> dict:
    pairs = [(to_number(r.get(x)), to_number(r.get(y))) for r in rows]
    pairs = [(a, b) for a, b in pairs if a is not None and b is not None]
    n = len(pairs)
    fit = None
    if n >= 3:
        xs = np.array([a for a, _ in pairs]); ys = np.array([b for _, b in pairs])
        if xs.std() > 0 and ys.std() > 0:
            slope, intercept = np.polyfit(xs, ys, 1)
            r = float(np.corrcoef(xs, ys)[0, 1])
            fit = {"slope": _r(float(slope), 6), "intercept": _r(float(intercept), 6), "r": _r(r, sig=False), "r2": _r(r * r, sig=False),
                   "x_min": _r(float(xs.min()), 6), "x_max": _r(float(xs.max()), 6)}
    kept = _thin(list(range(n)), limit)
    return {"x": x, "y": y, "n": n, "plotted": len(kept),
            "points": [{"x": _r(pairs[i][0], 6), "y": _r(pairs[i][1], 6)} for i in kept], "fit": fit}


def groups(rows: list[dict], by: str, value: str, limit: int = 30) -> dict:
    buckets: dict[str, list[float]] = {}
    for r in rows:
        v = to_number(r.get(value))
        if v is None:
            continue
        key = r.get(by)
        key = "(vide)" if is_empty(key) else str(key).strip()
        buckets.setdefault(key, []).append(v)
    ordered = sorted(buckets.items(), key=lambda kv: (-len(kv[1]), kv[0]))
    out = []
    for key, vals in ordered[:limit]:
        a = np.array(vals)
        out.append({"group": key, "n": int(a.size), "mean": _r(float(a.mean())),
                    "std": _r(float(a.std(ddof=1))) if a.size > 1 else None,
                    "min": _r(float(a.min())), "max": _r(float(a.max()))})
    # Quand tous les groupes tiennent dans le graphique : ordre naturel
    # (numérique pour des âges ou profondeurs, alphabétique sinon).
    if len(buckets) <= limit:
        if all(to_number(g["group"]) is not None for g in out):
            out.sort(key=lambda g: to_number(g["group"]))
        else:
            out.sort(key=lambda g: g["group"].casefold())
    return {"by": by, "value": value, "groups": out, "total_groups": len(buckets), "shown": len(out)}


_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Empêche qu'un texte soit interprété comme une formule par Excel (injection CSV)."""
    if isinstance(value, str):
        value = _CTRL.sub("", value)   # aussi pour les données déjà enregistrées
        if value.startswith(_FORMULA_START) and to_number(value) is None:
            return "'" + value
    return value


def export_rows(columns: list[str], rows: list[dict], numeric: Iterable[str]) -> list[list[Any]]:
    numeric = set(numeric)
    out = []
    for r in rows:
        line = []
        for c in columns:
            v = r.get(c)
            if c in numeric:
                n = to_number(v)
                line.append(n if n is not None else safe_cell(v))
            else:
                line.append(safe_cell(v))
        out.append(line)
    return out


def _fmt_number(x: float, decimal_comma: bool) -> str:
    s = repr(int(x)) if float(x).is_integer() and abs(x) < 1e15 else repr(x)
    return s.replace(".", ",") if decimal_comma else s


def to_csv(columns: list[str], rows: list[list[Any]], excel_fr: bool) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";" if excel_fr else ",", lineterminator="\r\n")
    w.writerow(columns)
    for r in rows:
        w.writerow(["" if v is None else _fmt_number(v, excel_fr) if isinstance(v, (int, float)) else v for v in r])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(title: str, columns: list[str], rows: list[list[Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "Données"
    ws.append([_CTRL.sub("", c) for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, c in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(c) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


DISTRIBUTIONS: dict[str, dict] = {
    "normal": {"label": "Normale", "positive": False,
               "use": "Mesures symétriques autour d'une moyenne (résistances, dimensions, teneurs)."},
    "lognormal": {"label": "Lognormale", "positive": True,
                  "use": "Valeurs positives asymétriques (perméabilités, résistances de sols, durées)."},
    "weibull": {"label": "Weibull", "positive": True,
                "use": "Ruptures et résistances de matériaux fragiles, durées de vie."},
    "gumbel": {"label": "Gumbel", "positive": False,
               "use": "Valeurs extrêmes : crues maximales annuelles, vents, charges maximales."},
    "exponential": {"label": "Exponentielle", "positive": True,
                    "use": "Temps entre événements (pannes, arrivées)."},
    "uniform": {"label": "Uniforme", "positive": False,
                "use": "Valeurs également probables entre un minimum et un maximum."},
}


class FitProblem(ValueError):
    """Erreur lisible renvoyée telle quelle à l'utilisateur."""


def _fit(name: str, x: np.ndarray):
    """Retourne (loi scipy figée, paramètres lisibles, nombre de paramètres libres)."""
    from scipy import stats as sps

    if name == "normal":
        mu, sigma = sps.norm.fit(x)
        return sps.norm(mu, sigma), [("μ (moyenne)", mu), ("σ (écart-type)", sigma)], 2
    if name == "lognormal":
        s, _, scale = sps.lognorm.fit(x, floc=0)
        return sps.lognorm(s, 0, scale), [("μ du logarithme", math.log(scale)), ("σ du logarithme", s)], 2
    if name == "weibull":
        c, _, scale = sps.weibull_min.fit(x, floc=0)
        return sps.weibull_min(c, 0, scale), [("k (forme)", c), ("λ (échelle)", scale)], 2
    if name == "gumbel":
        loc, scale = sps.gumbel_r.fit(x)
        return sps.gumbel_r(loc, scale), [("μ (position)", loc), ("β (échelle)", scale)], 2
    if name == "exponential":
        _, scale = sps.expon.fit(x, floc=0)
        return sps.expon(0, scale), [("λ (taux)", 1 / scale), ("Moyenne 1/λ", scale)], 1
    if name == "uniform":
        loc, scale = sps.uniform.fit(x)
        return sps.uniform(loc, scale), [("a (minimum)", loc), ("b (maximum)", loc + scale)], 2
    raise FitProblem(f"Loi inconnue : {name}.")


def fit_distributions(values: list[float], selected: str = "normal") -> dict:
    """Ajuste toutes les lois applicables, les classe (AIC) et détaille la loi choisie."""
    from scipy import stats as sps

    if selected not in DISTRIBUTIONS:
        raise FitProblem(f"Loi inconnue : {selected}.")
    x = np.asarray([v for v in values if v is not None and math.isfinite(v)], dtype=float)
    n = int(x.size)
    if n < 5:
        raise FitProblem("Il faut au moins 5 valeurs numériques pour ajuster une loi.")
    if float(x.std()) == 0.0:
        raise FitProblem("Toutes les valeurs sont identiques : aucune loi ne peut être ajustée.")
    positive = bool(np.all(x > 0))

    comparison = []
    frozen = {}
    for name, meta in DISTRIBUTIONS.items():
        entry = {"key": name, "label": meta["label"], "use": meta["use"]}
        if meta["positive"] and not positive:
            entry.update(available=False, reason="exige des valeurs strictement positives")
            comparison.append(entry)
            continue
        try:
            with np.errstate(all="ignore"):
                dist, params, k = _fit(name, x)
                loglik = float(np.sum(dist.logpdf(x)))
            if not math.isfinite(loglik) or not all(math.isfinite(v) for _, v in params):
                raise ValueError("ajustement non fini")
            ks = sps.kstest(x, dist.cdf)
            frozen[name] = (dist, params)
            entry.update(available=True, aic=_r(2 * k - 2 * loglik, 2, sig=False),
                         ks_d=_r(float(ks.statistic), 4, sig=False), ks_p=_r(float(ks.pvalue), 4, sig=False))
        except Exception:
            entry.update(available=False, reason="ajustement impossible sur ces données")
        comparison.append(entry)

    ranked = sorted((c for c in comparison if c["available"]), key=lambda c: c["aic"])
    for i, c in enumerate(ranked, start=1):
        c["rank"] = i
        c["delta_aic"] = _r(c["aic"] - ranked[0]["aic"], 2, sig=False)
    comparison.sort(key=lambda c: (not c["available"], c.get("rank", 99)))

    if selected not in frozen:
        reason = next(c.get("reason", "") for c in comparison if c["key"] == selected)
        raise FitProblem(f"La loi {DISTRIBUTIONS[selected]['label'].lower()} ne peut pas être utilisée ici : {reason}.")
    dist, params = frozen[selected]

    lo, hi = float(x.min()), float(x.max())
    pad = (hi - lo) * 0.15
    grid = np.linspace(lo - pad, hi + pad, 200)
    if DISTRIBUTIONS[selected]["positive"]:
        grid = grid[grid > 0]
    with np.errstate(all="ignore"):
        pdf = dist.pdf(grid)
    counts, edges = np.histogram(x, bins=int(min(30, max(5, math.ceil(math.sqrt(n))))))
    density = counts / (n * np.diff(edges))
    q = {p: float(dist.ppf(p)) for p in (0.05, 0.5, 0.95)}
    emp = {p: float(np.percentile(x, p * 100)) for p in (0.05, 0.5, 0.95)}

    sel = next(c for c in comparison if c["key"] == selected)
    return {
        "n": n,
        "summary": {"mean": _r(float(x.mean())), "std": _r(float(x.std(ddof=1))), "min": _r(lo), "max": _r(hi)},
        "selected": {"key": selected, "label": DISTRIBUTIONS[selected]["label"], "use": DISTRIBUTIONS[selected]["use"],
                     "parameters": [{"name": nm, "value": _r(float(v), 6)} for nm, v in params],
                     "aic": sel["aic"], "rank": sel["rank"], "ks_d": sel["ks_d"], "ks_p": sel["ks_p"]},
        "best": ranked[0]["key"],
        "quantiles": [{"p": p, "model": _r(q[p]), "empirical": _r(emp[p])} for p in (0.05, 0.5, 0.95)],
        "comparison": comparison,
        "plot": {"edges": [_r(float(e), 6) for e in edges], "density": [_r(float(d), 8) for d in density],
                 "x": [_r(float(v), 6) for v in grid],
                 "pdf": [(_r(float(v), 8) if math.isfinite(v) else None) for v in pdf]},
        "small_sample": n < 30,
    }


# ============================================================================
# SECTION 3 — BASE DE DONNÉES (SQLite sur PC, PostgreSQL en ligne)
# ============================================================================

def now_utc() -> datetime:
    return datetime.now(timezone.utc)


def ts(dt: datetime | None = None) -> str:
    """Horodatage texte, UTC, au format « AAAA-MM-JJ HH:MM:SS » sur les deux bases."""
    return (dt or now_utc()).strftime("%Y-%m-%d %H:%M:%S")


def ts_in(days: int) -> str:
    return ts(now_utc() + timedelta(days=days))


SQLITE_SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    full_name     TEXT NOT NULL,
    birth_date    TEXT,
    email         TEXT NOT NULL UNIQUE COLLATE NOCASE,
    password_hash TEXT NOT NULL,
    created_at    TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS user_sessions (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id    INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    token_hash TEXT NOT NULL UNIQUE,
    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    expires_at TEXT
);
CREATE TABLE IF NOT EXISTS datasets (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id      INTEGER REFERENCES users(id) ON DELETE CASCADE,
    name         TEXT NOT NULL,
    domain       TEXT NOT NULL,
    columns_json TEXT NOT NULL,
    created_at   TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    updated_at   TEXT
);
CREATE TABLE IF NOT EXISTS records (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    dataset_id  INTEGER NOT NULL REFERENCES datasets(id) ON DELETE CASCADE,
    values_json TEXT NOT NULL,
    created_at  TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
"""


_PG_NOW = "to_char(now() AT TIME ZONE 'UTC', 'YYYY-MM-DD HH24:MI:SS')"


POSTGRES_SCHEMA = f"""
CREATE TABLE IF NOT EXISTS users (
    id            BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
    full_name     TEXT NOT NULL,
    birth_date    TEXT,
    email         TEXT NOT NULL,
    password_hash TEXT NOT NULL,
    created_at    TEXT NOT NULL DEFAULT {_PG_NOW}
);
CREATE UNIQUE INDEX IF NOT EXISTS users_email_lower ON users (lower(email));
CREATE TABLE IF NOT EXISTS user_sessions (
    id         BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
    user_id    BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
    token_hash TEXT NOT NULL UNIQUE,
    created_at TEXT NOT NULL DEFAULT {_PG_NOW},
    expires_at TEXT
);
CREATE TABLE IF NOT EXISTS datasets (
    id           BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
    user_id      BIGINT REFERENCES users(id) ON DELETE CASCADE,
    name         TEXT NOT NULL,
    domain       TEXT NOT NULL,
    columns_json TEXT NOT NULL,
    created_at   TEXT NOT NULL DEFAULT {_PG_NOW},
    updated_at   TEXT
);
CREATE TABLE IF NOT EXISTS records (
    id          BIGINT GENERATED BY DEFAULT AS IDENTITY PRIMARY KEY,
    dataset_id  BIGINT NOT NULL REFERENCES datasets(id) ON DELETE CASCADE,
    values_json TEXT NOT NULL,
    created_at  TEXT NOT NULL DEFAULT {_PG_NOW}
);
"""


INDEXES = """
CREATE INDEX IF NOT EXISTS idx_sessions_user    ON user_sessions(user_id);
CREATE INDEX IF NOT EXISTS idx_sessions_expires ON user_sessions(expires_at);
CREATE INDEX IF NOT EXISTS idx_datasets_user    ON datasets(user_id);
CREATE INDEX IF NOT EXISTS idx_records_dataset  ON records(dataset_id, id);
"""


class Row(dict):
    """Ligne lisible par nom de colonne et par position."""

    __slots__ = ("_values",)

    def __init__(self, names: Sequence[str], values: Sequence[Any]):
        super().__init__(zip(names, values))
        self._values = tuple(values)

    def __getitem__(self, key):
        if isinstance(key, int):
            return self._values[key]
        return super().__getitem__(key)


class Cursor:
    def __init__(self, raw, kind: str):
        self._raw = raw
        self._kind = kind
        desc = raw.description
        self._names = [d[0] for d in desc] if desc else []

    def _wrap(self, values):
        if values is None:
            return None
        if self._kind == "sqlite":
            return values  # sqlite3.Row fait déjà les deux accès
        return Row(self._names, values)

    def fetchone(self):
        return self._wrap(self._raw.fetchone())

    def fetchall(self):
        return [self._wrap(v) for v in self._raw.fetchall()]

    def __iter__(self):
        return iter(self.fetchall())

    @property
    def rowcount(self) -> int:
        return self._raw.rowcount


_IS_NOT_PARAM = re.compile(r"\bIS NOT \?")


def _to_pg(sql: str) -> str:
    return _IS_NOT_PARAM.sub("IS DISTINCT FROM ?", sql).replace("?", "%s")


class Database:
    """Connexion unique pour une requête HTTP (ou une commande d'administration)."""

    def __init__(self, raw, kind: str, release=None):
        self.raw = raw
        self.kind = kind
        self._release = release

    def execute(self, sql: str, params: Sequence[Any] = ()) -> Cursor:
        if self.kind == "postgres":
            cur = self.raw.cursor()
            cur.execute(_to_pg(sql), params)
            return Cursor(cur, self.kind)
        return Cursor(self.raw.execute(sql, params), self.kind)

    def executemany(self, sql: str, seq) -> None:
        if self.kind == "postgres":
            with self.raw.cursor() as cur:
                cur.executemany(_to_pg(sql), list(seq))
        else:
            self.raw.executemany(sql, seq)

    def insert(self, sql: str, params: Sequence[Any] = ()) -> int:
        """INSERT qui renvoie l'identifiant de la ligne créée."""
        if self.kind == "postgres":
            cur = self.raw.cursor()
            cur.execute(_to_pg(sql) + " RETURNING id", params)
            return int(cur.fetchone()[0])
        return int(self.raw.execute(sql, params).lastrowid)

    def commit(self) -> None:
        self.raw.commit()

    def rollback(self) -> None:
        self.raw.rollback()

    def close(self) -> None:
        if self._release:
            self._release(self.raw)
        else:
            self.raw.close()


def _sqlite_columns(raw: sqlite3.Connection, table: str) -> dict[str, sqlite3.Row]:
    return {r["name"]: r for r in raw.execute(f"PRAGMA table_info({table})")}


def _migrate_sqlite(raw: sqlite3.Connection) -> None:
    """Met à niveau une base créée par la version 3 sans perdre de données."""
    users = _sqlite_columns(raw, "users")
    # v3 : birth_date était obligatoire. SQLite ne sait pas retirer un NOT NULL,
    # on reconstruit donc la table (procédure officielle « 12 étapes »).
    if users and users["birth_date"]["notnull"]:
        raw.execute("PRAGMA foreign_keys = OFF")
        raw.executescript(
            """
            BEGIN;
            CREATE TABLE users_new (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                full_name TEXT NOT NULL,
                birth_date TEXT,
                email TEXT NOT NULL UNIQUE COLLATE NOCASE,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            INSERT INTO users_new (id, full_name, birth_date, email, password_hash, created_at)
                SELECT id, full_name, NULLIF(birth_date, ''), email, password_hash,
                       COALESCE(created_at, CURRENT_TIMESTAMP) FROM users;
            DROP TABLE users;
            ALTER TABLE users_new RENAME TO users;
            COMMIT;
            """
        )
        raw.execute("PRAGMA foreign_keys = ON")

    sessions = _sqlite_columns(raw, "user_sessions")
    if "expires_at" not in sessions:
        raw.execute("ALTER TABLE user_sessions ADD COLUMN expires_at TEXT")
    # Les sessions v3 n'avaient pas d'expiration côté serveur : on les invalide.
    raw.execute("DELETE FROM user_sessions WHERE expires_at IS NULL")

    datasets = _sqlite_columns(raw, "datasets")
    if "user_id" not in datasets:  # base antérieure à la v3
        raw.execute("ALTER TABLE datasets ADD COLUMN user_id INTEGER")
    if "updated_at" not in datasets:
        raw.execute("ALTER TABLE datasets ADD COLUMN updated_at TEXT")
    raw.execute("UPDATE datasets SET updated_at = created_at WHERE updated_at IS NULL")

    # Enregistrements orphelins (jeu supprimé sans cascade dans les anciennes versions)
    raw.execute("DELETE FROM records WHERE dataset_id NOT IN (SELECT id FROM datasets)")
    raw.commit()


class Engine:
    """Une seule instance par serveur Streamlit (voir get_engine) : pool PostgreSQL ou chemin SQLite."""

    def __init__(self):
        self.pool = None
        if USE_POSTGRES:
            from psycopg_pool import ConnectionPool
            self.pool = ConnectionPool(
                DATABASE_URL, min_size=0, max_size=5, open=True, timeout=30,
                check=ConnectionPool.check_connection,
                kwargs={"autocommit": False, "connect_timeout": 15, "prepare_threshold": None},
            )
        self.init_db()

    def connect(self) -> Database:
        if self.pool is not None:
            raw = self.pool.getconn()

            def release(c):
                try:
                    if not c.closed:
                        c.rollback()
                finally:
                    self.pool.putconn(c)
            return Database(raw, "postgres", release)
        raw = sqlite3.connect(DB_PATH, timeout=10, check_same_thread=False)
        raw.row_factory = sqlite3.Row
        raw.execute("PRAGMA foreign_keys = ON")
        return Database(raw, "sqlite")

    def init_db(self) -> None:
        if self.pool is not None:
            db = self.connect()
            try:
                with db.raw.cursor() as cur:
                    cur.execute("SELECT pg_advisory_xact_lock(4242001)")
                    for stmt in (POSTGRES_SCHEMA + INDEXES).split(";"):
                        if stmt.strip():
                            cur.execute(stmt)
                db.commit()
            finally:
                db.close()
            return
        DB_PATH.parent.mkdir(parents=True, exist_ok=True)
        db = self.connect()
        try:
            db.raw.execute("PRAGMA journal_mode = WAL")
            db.raw.executescript(SQLITE_SCHEMA)
            _migrate_sqlite(db.raw)
            db.raw.executescript(INDEXES)
            db.raw.commit()
        finally:
            db.close()


@st.cache_resource(show_spinner="Connexion à la base de données…")
def get_engine() -> Engine:
    # Streamlit réexécute le script à chaque clic : sans ce cache, chaque clic
    # ouvrirait un nouveau pool de connexions.
    return Engine()


@contextmanager
def database():
    db = get_engine().connect()
    try:
        yield db
    finally:
        db.close()


def _integrity_errors() -> tuple[type[BaseException], ...]:
    errors: list[type[BaseException]] = [sqlite3.IntegrityError]
    try:
        import psycopg
        errors.append(psycopg.errors.IntegrityError)
    except ImportError:
        pass
    return tuple(errors)


IntegrityError = _integrity_errors()

# ============================================================================
# SECTION 4 — ESPACES PERSONNELS (SANS COMPTE)
# ============================================================================
#
# Pas d'inscription : le premier passage crée un jeton aléatoire, placé dans
# l'adresse (?espace=…). Le jeton n'est jamais stocké en clair : la base garde
# seulement son empreinte SHA-256. L'espace n'est créé en base qu'au premier
# jeu enregistré (une simple visite ne laisse aucune trace).

TOKEN_RE = re.compile(r"^[A-Za-z0-9_-]{20,64}$")


def _space_email(token: str) -> str:
    return "espace-" + hashlib.sha256(token.encode("utf-8")).hexdigest()[:48] + SPACE_DOMAIN


def current_token() -> str:
    token = str(st.query_params.get(SPACE_PARAM, "") or "")
    if not TOKEN_RE.match(token):
        token = st.session_state.get("_space_token") or secrets.token_urlsafe(18)
        st.query_params[SPACE_PARAM] = token
    st.session_state["_space_token"] = token
    return token


def find_space(token: str) -> int | None:
    with database() as db:
        row = db.execute("SELECT id FROM users WHERE email = ?", (_space_email(token),)).fetchone()
    return row["id"] if row else None


def ensure_space(user: dict) -> int:
    """Crée l'espace en base au premier enregistrement et renvoie son identifiant."""
    if user.get("id"):
        return user["id"]
    email = _space_email(user["token"])
    with database() as db:
        try:
            uid = db.insert("INSERT INTO users (full_name, email, password_hash, created_at) VALUES (?, ?, ?, ?)",
                            ("Espace personnel", email, "!", ts()))
            db.commit()
        except IntegrityError:  # créé au même instant dans un autre onglet
            db.rollback()
            uid = db.execute("SELECT id FROM users WHERE email = ?", (email,)).fetchone()["id"]
    user["id"] = uid
    return uid


def space_link(token: str) -> str:
    base = PUBLIC_URL
    try:
        url = st.context.url
        if url and "/~/+" not in url and "localhost" in url:
            base = url.split("?")[0]
    except Exception:
        pass
    return base.rstrip("/") + f"/?{SPACE_PARAM}={token}"


def new_space() -> None:
    token = secrets.token_urlsafe(18)
    st.session_state["_space_token"] = token
    st.query_params[SPACE_PARAM] = token
    flash("Nouvel espace vide créé. Gardez son lien : il est dans « Mon espace ».")
    go("home", ds_id=None)


def wipe_space(user_id: int | None) -> int:
    """Supprime tous les jeux de l'espace (l'espace lui-même reste)."""
    if not user_id:
        return 0
    with database() as db:
        ids = [r[0] for r in db.execute("SELECT id FROM datasets WHERE user_id = ?", (user_id,))]
        db.executemany("DELETE FROM records WHERE dataset_id = ?", [(i,) for i in ids])
        db.execute("DELETE FROM datasets WHERE user_id = ?", (user_id,))
        db.commit()
    return len(ids)


@st.cache_resource(ttl=24 * 3600, show_spinner=False)
def purge_idle_spaces() -> int:
    """Une fois par jour : supprime les espaces personnels inutilisés depuis SPACE_IDLE_DAYS jours."""
    cutoff = ts(now_utc() - timedelta(days=SPACE_IDLE_DAYS))
    with database() as db:
        rows = db.execute(
            """SELECT u.id FROM users u
               WHERE u.email LIKE ? AND u.created_at < ?
                 AND NOT EXISTS (SELECT 1 FROM datasets d WHERE d.user_id = u.id
                                 AND COALESCE(d.updated_at, d.created_at) >= ?)""",
            ("%" + SPACE_DOMAIN, cutoff, cutoff)).fetchall()
        ids = [r[0] for r in rows]
        for uid in ids:
            dids = [r[0] for r in db.execute("SELECT id FROM datasets WHERE user_id = ?", (uid,))]
            db.executemany("DELETE FROM records WHERE dataset_id = ?", [(i,) for i in dids])
            db.execute("DELETE FROM datasets WHERE user_id = ?", (uid,))
            db.execute("DELETE FROM users WHERE id = ?", (uid,))
        db.commit()
    return len(ids)


# ============================================================================
# SECTION 5 — OPÉRATIONS SUR LES DONNÉES
# ============================================================================

class AppError(Exception):
    """Erreur lisible affichée telle quelle à l'utilisateur."""


EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]{2,}$")


def _dataset_out(row) -> dict:
    return {"id": row["id"], "name": row["name"], "domain": row["domain"],
            "columns": json.loads(row["columns_json"]), "record_count": row["record_count"],
            "created_at": row["created_at"], "updated_at": row["updated_at"] or row["created_at"]}


_DS_SELECT = """SELECT d.*, (SELECT COUNT(*) FROM records r WHERE r.dataset_id = d.id) AS record_count
                FROM datasets d"""


def list_datasets(user_id: int | None) -> list[dict]:
    """Les jeux de cet espace uniquement."""
    if not user_id:
        return []
    with database() as db:
        rows = db.execute(_DS_SELECT + " WHERE d.user_id = ? ORDER BY COALESCE(d.updated_at, d.created_at) DESC, d.id DESC",
                          (user_id,)).fetchall()
    return [_dataset_out(r) for r in rows]


def get_dataset(db: Database, did: int, user_id: int) -> dict:
    row = db.execute(_DS_SELECT + " WHERE d.id = ? AND d.user_id = ?", (did, user_id)).fetchone()
    if not row:
        raise AppError("Ce jeu de données n'existe pas dans votre espace (il a peut-être été supprimé).")
    return _dataset_out(row)


def check_quota(db: Database, user_id: int, new_datasets: int = 0, new_records: int = 0) -> None:
    row = db.execute(
        """SELECT (SELECT COUNT(*) FROM datasets WHERE user_id = ?) AS n_ds,
                  (SELECT COUNT(*) FROM records r JOIN datasets d ON d.id = r.dataset_id WHERE d.user_id = ?) AS n_rec,
                  (SELECT COUNT(*) FROM records) AS n_all""", (user_id, user_id)).fetchone()
    if new_records and row["n_all"] + new_records > MAX_RECORDS_TOTAL:
        raise AppError("La base de calculBA est pleine pour le moment. Supprimez des jeux inutiles ou réessayez plus tard.")
    if new_datasets and row["n_ds"] + new_datasets > MAX_DATASETS_PER_USER:
        raise AppError(f"Limite de {MAX_DATASETS_PER_USER} jeux par espace atteinte. Supprimez des jeux inutiles.")
    if new_records and row["n_rec"] + new_records > MAX_RECORDS_PER_USER:
        left = max(0, MAX_RECORDS_PER_USER - row["n_rec"])
        raise AppError(f"Limite de {MAX_RECORDS_PER_USER:,} observations par espace : il reste de la place pour "
                       f"{left:,}.".replace(",", " "))


def _insert_dataset(db: Database, user_id: int, name: str, domain: str, columns: list[str]) -> int:
    now = ts()
    return db.insert(
        "INSERT INTO datasets (user_id, name, domain, columns_json, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
        (user_id, name, domain if domain in DOMAINS else "Autre", json.dumps(columns, ensure_ascii=False), now, now))


def create_dataset(user_id: int, name: str, domain: str, columns_text: str) -> int:
    name = name.strip()
    if not name:
        raise AppError("Donnez un nom au jeu de données.")
    if len(name) > MAX_NAME_LENGTH:
        raise AppError(f"Nom trop long ({MAX_NAME_LENGTH} caractères maximum).")
    columns = [c.strip() for c in re.split(r"[\n,;\t]+", columns_text) if c.strip()]
    try:
        columns = validate_columns(columns)
    except ImportProblem as e:
        raise AppError(str(e))
    with database() as db:
        check_quota(db, user_id, new_datasets=1)
        did = _insert_dataset(db, user_id, name, domain, columns)
        db.commit()
    return did


def import_dataset(user_id: int, filename: str, raw: bytes, name: str, domain: str) -> tuple[int, int]:
    if len(raw) > MAX_UPLOAD_BYTES:
        raise AppError(f"Fichier trop volumineux (maximum {MAX_UPLOAD_BYTES // (1024 * 1024)} Mo).")
    try:
        columns, data = read_table(filename, raw)
    except ImportProblem as e:
        raise AppError(str(e))
    ds_name = (name.strip() or Path(filename).stem or "Données importées")[:MAX_NAME_LENGTH]
    with database() as db:
        check_quota(db, user_id, new_datasets=1, new_records=len(data))
        did = _insert_dataset(db, user_id, ds_name, domain, columns)
        now = ts()
        db.executemany("INSERT INTO records (dataset_id, values_json, created_at) VALUES (?, ?, ?)",
                       [(did, json.dumps(dict(zip(columns, row)), ensure_ascii=False), now) for row in data])
        db.commit()
    return did, len(data)


def update_dataset(did: int, user_id: int, name: str, domain: str) -> None:
    name = name.strip()
    if not name:
        raise AppError("Donnez un nom au jeu de données.")
    with database() as db:
        get_dataset(db, did, user_id)
        db.execute("UPDATE datasets SET name = ?, domain = ?, updated_at = ? WHERE id = ?",
                   (name[:MAX_NAME_LENGTH], domain if domain in DOMAINS else "Autre", ts(), did))
        db.commit()


def delete_dataset(did: int, user_id: int) -> None:
    with database() as db:
        get_dataset(db, did, user_id)
        db.execute("DELETE FROM records WHERE dataset_id = ?", (did,))
        db.execute("DELETE FROM datasets WHERE id = ? AND user_id = ?", (did, user_id))
        db.commit()


def load_records(db: Database, did: int) -> list[tuple[int, dict]]:
    """(identifiant, valeurs) de chaque observation, dans l'ordre de saisie."""
    return [(r["id"], json.loads(r["values_json"]))
            for r in db.execute("SELECT id, values_json FROM records WHERE dataset_id = ? ORDER BY id", (did,))]


def _clean_record(db: Database, did: int, ds: dict, values: dict, skip_rid: int | None = None) -> dict:
    clean = {c: clean_value(values.get(c)) for c in ds["columns"]}
    if all(v is None for v in clean.values()):
        raise AppError("L'observation est vide : renseignez au moins une valeur.")
    # Texte refusé dans une colonne numérique (ex. « abc » dans Resistance_MPa).
    suspects = [c for c in ds["columns"] if clean[c] is not None and to_number(clean[c]) is None]
    if suspects:
        rows = [json.loads(r["values_json"]) for r in db.execute(
            "SELECT values_json FROM records WHERE dataset_id = ? AND id IS NOT ?", (did, skip_rid))]
        prof = profile(suspects, rows)
        bad = [c for c in suspects if prof[c]["numeric"]]
        if bad:
            raise AppError(f"Valeur non numérique pour : {', '.join(bad)}. Corrigez-la (ex. 30,5) ou laissez le champ vide.")
    return clean


def add_record(did: int, user_id: int, values: dict) -> None:
    with database() as db:
        ds = get_dataset(db, did, user_id)
        if ds["record_count"] >= MAX_ROWS:
            raise AppError(f"Ce jeu a atteint la limite de {MAX_ROWS} observations.")
        check_quota(db, user_id, new_records=1)
        clean = _clean_record(db, did, ds, values)
        db.insert("INSERT INTO records (dataset_id, values_json, created_at) VALUES (?, ?, ?)",
                  (did, json.dumps(clean, ensure_ascii=False), ts()))
        db.execute("UPDATE datasets SET updated_at = ? WHERE id = ?", (ts(), did))
        db.commit()


def update_record(did: int, user_id: int, rid: int, values: dict) -> None:
    with database() as db:
        ds = get_dataset(db, did, user_id)
        clean = _clean_record(db, did, ds, values, skip_rid=rid)
        cur = db.execute("UPDATE records SET values_json = ? WHERE id = ? AND dataset_id = ?",
                         (json.dumps(clean, ensure_ascii=False), rid, did))
        if cur.rowcount == 0:
            raise AppError("Observation introuvable.")
        db.execute("UPDATE datasets SET updated_at = ? WHERE id = ?", (ts(), did))
        db.commit()


def delete_record(did: int, user_id: int, rid: int) -> None:
    with database() as db:
        get_dataset(db, did, user_id)
        cur = db.execute("DELETE FROM records WHERE id = ? AND dataset_id = ?", (rid, did))
        if cur.rowcount == 0:
            raise AppError("Observation introuvable.")
        db.execute("UPDATE datasets SET updated_at = ? WHERE id = ?", (ts(), did))
        db.commit()


def export_bytes(did: int, user_id: int, fmt: str) -> bytes:
    with database() as db:
        ds = get_dataset(db, did, user_id)
        rows = [v for _, v in load_records(db, did)]
    numeric_cols = [c for c, p in profile(ds["columns"], rows).items() if p["numeric"]]
    table = export_rows(ds["columns"], rows, numeric_cols)
    if fmt == "xlsx":
        return to_xlsx(ds["name"], ds["columns"], table)
    return to_csv(ds["columns"], table, excel_fr=(fmt == "csv-excel"))


# ============================================================================
# SECTION 6 — INTERFACE STREAMLIT
# ============================================================================

def fmt(v, digits: int = 4) -> str:
    """Nombre à la française : 26,9471 ; 1 234,5 ; petites valeurs en chiffres significatifs."""
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return "—"
    if isinstance(v, (int, np.integer)) or (isinstance(v, float) and v.is_integer() and abs(v) < 1e15):
        return f"{int(v):,}".replace(",", " ")
    if v != 0 and abs(v) < 0.01:
        return f"{v:.4g}".replace(".", ",")
    s = f"{v:,.{digits}f}".rstrip("0").rstrip(".")
    return s.replace(",", " ").replace(".", ",")


def plural(n: int, one: str, many: str) -> str:
    return f"{fmt(n)} {one if n <= 1 else many}"


def flash(message: str, icon: str | None = None) -> None:
    st.session_state.setdefault("flash", []).append((message, icon))


def show_flash() -> None:
    for message, icon in st.session_state.pop("flash", []):
        st.toast(message, icon=icon)


def go(view: str, **state) -> None:
    """Change de page au prochain passage du script.

    Streamlit interdit de modifier un sélecteur déjà affiché pendant le passage
    en cours : la demande est mise en attente et appliquée au début du suivant.
    """
    st.session_state["_pending"] = {"view": view, **state}
    st.rerun()


def apply_pending() -> None:
    for key, value in st.session_state.pop("_pending", {}).items():
        if value is None:
            st.session_state.pop(key, None)
        else:
            st.session_state[key] = value


# ------------------------------------------------------------ apparence
#
# Style « carnet d'ingénieur » : papier clair, encre noire, une seule couleur
# d'accent (vert), titres à empattements, chiffres en police technique.

DOMAIN_COLOR = {
    "Béton": "#3B5B7A", "Sols et géotechnique": "#8A5A2B", "Topographie": "#3F6B3A", "Matériaux": "#9A6B1F",
    "Structures": "#5B4A7A", "Routes": "#4A4A46", "Hydraulique": "#2D6B78", "Autre": "#6E685D",
}


def domain_color(domain: str) -> str:
    return DOMAIN_COLOR.get(domain, DOMAIN_COLOR["Autre"])


def inject_css() -> None:
    st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Serif:wght@500;600&family=IBM+Plex+Mono:wght@500&display=swap');
:root { --paper:#F6F4EF; --side:#EEEAE1; --ink:#1B1A17; --muted:#6E685D; --rule:#D9D3C7; --card:#FFFFFF; --accent:#1F5C4A; }
.stApp, .stApp p, .stApp li, .stApp label, .stApp input, .stApp textarea, .stApp button, .stApp td, .stApp th,
.stApp span:not([data-testid="stIconMaterial"]) { font-family: 'IBM Plex Sans', system-ui, sans-serif; }
[data-testid="stIconMaterial"] { font-family: 'Material Symbols Rounded' !important; }
.stApp { background: var(--paper); color: var(--ink); }
.stApp h1, .stApp h2, .stApp h3 { font-family: 'IBM Plex Serif', Georgia, serif !important; font-weight: 600; color: var(--ink); letter-spacing: -.01em; }
header[data-testid="stHeader"] { background: transparent; }
.block-container { padding-top: 2.4rem; max-width: 1160px; }

[data-testid="stSidebar"] { background: var(--side); border-right: 1px solid var(--rule); }
[data-testid="stSidebar"] .stButton button { border-radius: 0; justify-content: flex-start; box-shadow: none;
  border: none; border-left: 3px solid transparent; background: transparent; color: var(--ink); padding: 6px 12px; min-height: 0; }
[data-testid="stSidebar"] .stButton button > div { justify-content: flex-start; width: 100%; }
[data-testid="stSidebar"] .stButton button p { text-align: left; font-size: .95rem; }
[data-testid="stSidebar"] .stButton button:hover { background: rgba(27,26,23,.05); color: var(--ink); }
.stApp [data-testid="stSidebar"] .stButton > button[kind="primary"] { border: none; border-left: 3px solid var(--accent);
  background: rgba(31,92,74,.09); box-shadow: none; }
.stApp [data-testid="stSidebar"] .stButton > button[kind="primary"]:hover { background: rgba(31,92,74,.14); }
.stApp [data-testid="stSidebar"] .stButton > button[kind="primary"] p { color: var(--accent) !important; font-weight: 600; }

.stApp .stButton > button, .stApp .stDownloadButton > button, .stApp .stFormSubmitButton > button {
  border-radius: 4px; font-weight: 500; box-shadow: none; }
.stApp .stButton > button[kind="primary"], .stApp .stFormSubmitButton > button[kind="primaryFormSubmit"] {
  background: var(--accent); border: 1px solid var(--accent); color: #fff; }
.stApp .stButton > button[kind="primary"]:hover { background: #17473A; border-color: #17473A; }
.stApp .stButton > button[kind="secondary"], .stApp .stDownloadButton > button {
  background: var(--card); border: 1px solid var(--rule); color: var(--ink); }
.stApp .stButton > button[kind="secondary"]:hover, .stApp .stDownloadButton > button:hover { border-color: var(--ink); color: var(--ink); }
.stApp [data-testid="stExpander"] details { background: var(--card); border: 1px solid var(--rule); border-radius: 6px; }
.stApp [data-testid="stForm"] { background: var(--card); border: 1px solid var(--rule); border-radius: 6px; }
.stApp [data-testid="stDataFrame"] { border: 1px solid var(--rule); border-radius: 6px; }
.stApp [data-testid="stAlertContainer"] { border-radius: 4px; }
.stApp [data-testid="stCaptionContainer"] { color: var(--muted); }

.cb-brand { padding: 4px 0 14px; border-bottom: 1px solid var(--rule); margin-bottom: 14px; }
.cb-brand b { font-family: 'IBM Plex Serif', Georgia, serif; font-size: 1.55rem; font-weight: 600; color: var(--ink); letter-spacing: -.02em; }
.cb-brand b span { color: var(--accent); }
.cb-brand small { display: block; font-family: 'IBM Plex Mono', monospace; font-size: .68rem; letter-spacing: .14em;
  text-transform: uppercase; color: var(--muted); margin-top: 2px; }
.cb-label { font-family: 'IBM Plex Mono', monospace; font-size: .7rem; letter-spacing: .14em; text-transform: uppercase; color: var(--muted); }
.cb-mast { border-bottom: 1px solid var(--ink); padding-bottom: 16px; margin-bottom: 20px; }
.cb-mast h1 { font-size: 2.3rem; line-height: 1.15; margin: 6px 0 8px !important; padding: 0 !important; max-width: 24ch; }
.cb-mast p { color: var(--muted); max-width: 68ch; margin: 0; font-size: 1.02rem; }
.cb-stats { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); border-top: 1px solid var(--rule);
  border-bottom: 1px solid var(--rule); margin: 0 0 22px; }
.cb-stat { padding: 12px 16px 14px; border-left: 1px solid var(--rule); }
.cb-stat:first-child { border-left: none; padding-left: 0; }
.cb-stat .v { font-family: 'IBM Plex Serif', Georgia, serif; font-size: 1.75rem; font-weight: 600; color: var(--ink); margin-top: 2px; }
.cb-h2 { font-family: 'IBM Plex Serif', Georgia, serif; font-weight: 600; font-size: 1.3rem; color: var(--ink); margin: 26px 0 4px; }
.cb-sub { color: var(--muted); margin: 0 0 14px; }
.cb-steps { display: grid; grid-template-columns: repeat(auto-fit, minmax(220px, 1fr)); gap: 0; border-top: 1px solid var(--rule); margin-bottom: 18px; }
.cb-step { padding: 14px 18px 16px 0; }
.cb-step + .cb-step { padding-left: 18px; border-left: 1px solid var(--rule); }
.cb-step .n { font-family: 'IBM Plex Mono', monospace; color: var(--accent); font-size: .8rem; }
.cb-step b { display: block; font-family: 'IBM Plex Serif', Georgia, serif; font-size: 1.1rem; margin: 4px 0; color: var(--ink); }
.cb-step span { color: var(--muted); font-size: .92rem; }
.cb-card-title { font-family: 'IBM Plex Serif', Georgia, serif; font-weight: 600; font-size: 1.12rem; color: var(--ink); margin: 4px 0 6px; line-height: 1.3; }
.cb-desc { color: var(--muted); font-size: .9rem; margin: 0 0 10px; min-height: 2.7em; }
[class*="st-key-cbcard"] { background: var(--card); border: 1px solid var(--rule) !important; border-radius: 6px !important; }
[class*="st-key-cbcard"]:hover { border-color: #B9B1A2 !important; }
</style>
""", unsafe_allow_html=True)


def masthead(label: str, title: str, lead: str = "", color: str = "") -> None:
    style = f' style="color:{color}"' if color else ""
    st.markdown(f'<div class="cb-mast"><div class="cb-label"{style}>{html.escape(label)}</div>'
                f'<h1>{html.escape(title)}</h1>{f"<p>{html.escape(lead)}</p>" if lead else ""}</div>',
                unsafe_allow_html=True)


def stat_strip(items: list[tuple[str, str]]) -> None:
    cells = "".join(f'<div class="cb-stat"><div class="cb-label">{html.escape(a)}</div><div class="v">{html.escape(b)}</div></div>'
                    for a, b in items)
    st.markdown(f'<div class="cb-stats">{cells}</div>', unsafe_allow_html=True)


def card_title(domain: str, title: str) -> None:
    st.markdown(f'<div class="cb-label" style="color:{domain_color(domain)}">{html.escape(domain)}</div>'
                f'<div class="cb-card-title">{html.escape(title)}</div>', unsafe_allow_html=True)


# ------------------------------------------------------------ exemples par domaine

def _csv(header: list[str], rows: list[list]) -> str:
    def cell(v):
        if isinstance(v, float):
            return f"{v}".replace(".", ",")
        return "" if v is None else str(v)
    return "\n".join([";".join(header)] + [";".join(cell(v) for v in r) for r in rows]) + "\n"


def _example_sols() -> str:
    rng = np.random.default_rng(11)
    rows = []
    for s in range(1, 5):
        for k in range(1, 13):
            z = k * 0.5
            nature = "Argile latéritique" if z <= 2.5 else ("Argile sableuse" if z <= 4.5 else "Sable argileux")
            w = round(float(rng.normal(19 + 1.6 * z, 1.5)), 1)
            gd = round(float(rng.normal(1.82 - 0.035 * z, 0.03)), 2)
            ll = round(float(rng.normal(52 - 2.5 * z, 3)), 0)
            ip = round(float(rng.normal(24 - 1.4 * z, 2)), 0)
            cbr = round(float(rng.normal(28 - 6 * z, 3)), 0) if z <= 2.0 else None
            rows.append([f"SP{s}", float(z), nature, w, gd, ll, ip, cbr])
    return _csv(["Sondage", "Profondeur_m", "Nature", "Teneur_eau_pct", "Densite_seche_t_m3",
                 "Limite_liquidite_pct", "Indice_plasticite_pct", "CBR_pct"], rows)


def _example_topo() -> str:
    rng = np.random.default_rng(7)
    rows = []
    n = 0
    for i in range(6):
        for j in range(7):
            n += 1
            e, nn = 774100 + 25 * i + float(rng.normal(0, 2)), 428200 + 25 * j + float(rng.normal(0, 2))
            alt_m = 742.0 + 0.045 * (e - 774100) - 0.03 * (nn - 428200) + float(rng.normal(0, .25))
            code = "Repère" if n in (1, 21, 42) else ("Axe route" if j == 3 else "Terrain naturel")
            rows.append([f"P{n:02d}", round(e, 3), round(nn, 3), round(alt_m, 3), code])
    return _csv(["Point", "Est_m", "Nord_m", "Altitude_m", "Code"], rows)


def _example_granulo() -> str:
    sieves = [0.08, 0.16, 0.315, 0.63, 1.25, 2.5, 4.0, 5.0, 6.3, 8.0, 10.0, 12.5, 16.0, 20.0, 25.0]
    mats = {"Sable 0/4": (0.9, 1.6), "Gravillon 4/10": (6.5, 4.0), "Gravier 10/20": (14.0, 5.5)}
    rows = []
    for name, (d50, slope) in mats.items():
        for d in sieves:
            passing = 100 / (1 + (d50 / d) ** slope)
            passing = 100.0 if passing > 99.5 else round(passing, 1)
            rows.append([name, d, passing, round(100 - passing, 1)])
    return _csv(["Echantillon", "Tamis_mm", "Passant_pct", "Refus_cumule_pct"], rows)


def _example_structures() -> str:
    rng = np.random.default_rng(5)
    rows = []
    e_mod, inertia = 30e9, 0.25 * 0.5 ** 3 / 12  # béton 30 GPa, poutre 25 × 50 cm
    for k in range(1, 31):
        level = ["RDC", "R+1", "R+2"][(k - 1) // 10]
        span = round(float(rng.uniform(3.5, 7.5)), 2)
        q = round(float(rng.uniform(15, 45)), 1)
        f_inst = 5 * q * 1e3 * span ** 4 / (384 * e_mod * inertia) * 1000  # mm
        f = round(f_inst * float(rng.uniform(1.6, 2.8)), 1)  # flèche à long terme (fluage, fissuration)
        f_adm = round(span * 1000 / 250, 1)
        rows.append([f"P{k}", level, span, q, f, f_adm, "Oui" if f <= f_adm else "Non"])
    return _csv(["Element", "Niveau", "Portee_m", "Charge_kN_m", "Fleche_mm", "Fleche_admissible_mm", "Conforme"], rows)


def _example_routes() -> str:
    rng = np.random.default_rng(3)
    rows = []
    for k in range(25):
        pr = round(k * 0.2, 1)
        section = "A" if pr < 2.5 else "B"
        ep = round(float(rng.normal(20 if section == "A" else 17, 1.2)), 1)
        defl = round(float(rng.normal(55 if section == "A" else 78, 9)), 0)
        comp = round(float(rng.normal(98.5, 1.2)), 1)
        w = round(float(rng.normal(7.5, .8)), 1)
        rows.append([pr, section, ep, defl, comp, w])
    return _csv(["PR_km", "Section", "Epaisseur_base_cm", "Deflexion_1_100mm", "Compacite_pct", "Teneur_eau_pct"], rows)


def _example_hydro() -> str:
    rng = np.random.default_rng(9)
    # Régime équatorial : deux saisons des pluies (mai et octobre), étiage en février.
    shape = [0.35, 0.25, 0.45, 0.8, 1.0, 0.85, 0.6, 0.65, 1.1, 1.6, 1.3, 0.6]
    rows = []
    for station, base, a, b in (("Station A", 180.0, 0.32, 0.45), ("Station B", 95.0, 0.41, 0.42)):
        for year in (2024, 2025):
            for m in range(12):
                q = round(base * shape[m] * float(rng.normal(1, .1)), 1)
                h = round(a * q ** b, 2)
                rows.append([station, f"{year}-{m + 1:02d}", q, h])
    return _csv(["Station", "Mois", "Debit_m3_s", "Hauteur_m"], rows)


EXAMPLE_SETS = {
    "Béton": ("Essais de compression", "54 éprouvettes de 3 lots à 7, 14 et 28 jours : résistance et affaissement.",
              lambda: EXAMPLE_CSV),
    "Sols et géotechnique": ("Sondages de sol", "4 sondages tous les 0,5 m : teneur en eau, densité sèche, limites d'Atterberg, CBR.",
                             _example_sols),
    "Topographie": ("Levé topographique", "42 points en coordonnées (Est, Nord, Altitude) avec codes de levé.",
                    _example_topo),
    "Matériaux": ("Analyse granulométrique", "Courbes de 3 granulats (sable 0/4, gravillon 4/10, gravier 10/20).",
                  _example_granulo),
    "Structures": ("Flèches de poutres", "30 poutres : portée, charge, flèche mesurée et admissible (L/250).",
                   _example_structures),
    "Routes": ("Contrôle de chaussée", "25 points tous les 200 m : épaisseur, déflexion, compacité, teneur en eau.",
               _example_routes),
    "Hydraulique": ("Débits mensuels", "2 stations sur 2 ans : débit et hauteur d'eau, avec les saisons des pluies.",
                    _example_hydro),
}


def create_example(user: dict, domain: str) -> None:
    title, _, make = EXAMPLE_SETS[domain]
    try:
        uid = ensure_space(user)
        did, n = import_dataset(uid, f"{title}.csv", make().encode("utf-8-sig"), f"Exemple — {title}", domain)
        flash(f"Exemple « {title} » créé : {plural(n, 'ligne', 'lignes')}.")
        go("dataset", ds_id=did, tab="Statistiques", edit_row=None)
    except AppError as e:
        st.error(str(e))


def examples_gallery(user: dict, key: str) -> None:
    domains = list(EXAMPLE_SETS)
    for start in range(0, len(domains), 3):
        cols = st.columns(3)
        for col, domain in zip(cols, domains[start:start + 3]):
            title, desc, make = EXAMPLE_SETS[domain]
            with col.container(border=True, key=f"cbcard_{key}_{domain.split()[0]}"):
                card_title(domain, title)
                st.markdown(f'<div class="cb-desc">{html.escape(desc)}</div>', unsafe_allow_html=True)
                b1, b2 = st.columns(2)
                if b1.button("Utiliser", key=f"{key}_try_{domain}", type="primary", width="stretch"):
                    create_example(user, domain)
                b2.download_button("Fichier CSV", make().encode("utf-8-sig"), file_name=f"exemple_{domain.split()[0].lower()}.csv",
                                   mime="text/csv", key=f"{key}_dl_{domain}", width="stretch")
    st.caption("Données fictives mais réalistes, faites pour apprendre à utiliser calculBA.")


# ------------------------------------------------------------ navigation

def nav_button(label: str, view: str, **state) -> None:
    active = st.session_state.get("view", "home") == view
    if st.button(label, width="stretch", type="primary" if active else "tertiary", key=f"nav_{view}"):
        go(view, **state)


def sidebar(user: dict, datasets: list[dict]) -> None:
    with st.sidebar:
        st.markdown('<div class="cb-brand"><b>calcul<span>BA</span></b><small>Données de génie civil</small></div>',
                    unsafe_allow_html=True)
        nav_button("Mes jeux de données", "home")
        nav_button("Nouveau jeu", "new")
        nav_button("Exemples par domaine", "examples")
        nav_button("Mode d'emploi", "guide")
        if datasets:
            st.markdown('<div class="cb-label" style="margin:18px 0 4px">Récents</div>', unsafe_allow_html=True)
            for d in datasets[:8]:
                current = st.session_state.get("view") == "dataset" and st.session_state.get("ds_id") == d["id"]
                if st.button(d["name"], key=f"recent_{d['id']}", width="stretch",
                             type="primary" if current else "tertiary"):
                    go("dataset", ds_id=d["id"], tab="Données", edit_row=None)
        st.divider()
        my_space_box(user, datasets)
        if not USE_POSTGRES and str(ROOT).startswith("/mount/src"):  # Streamlit Community Cloud sans base Neon
            st.warning("Base temporaire : sur Streamlit Cloud, les données sont effacées au redémarrage. "
                       "Ajoutez DATABASE_URL (base Neon) dans les Secrets.")


def my_space_box(user: dict, datasets: list[dict]) -> None:
    with st.expander("Mon espace", expanded=False):
        st.caption("Votre travail est privé et lié à ce lien personnel. "
                   "Gardez-le (favori, note) pour le retrouver sur n'importe quel appareil.")
        st.code(space_link(user["token"]), language=None, wrap_lines=True)
        if datasets:
            sure = st.checkbox(f"Effacer définitivement mes {plural(len(datasets), 'jeu', 'jeux')}", key="wipe_sure")
            if st.button("Tout effacer", disabled=not sure, width="stretch", key="wipe_btn"):
                n = wipe_space(user["id"])
                st.session_state.pop("wipe_sure", None)
                flash(f"Espace vidé : {plural(n, 'jeu supprimé', 'jeux supprimés')}.")
                go("home", ds_id=None)
        if st.button("Nouvel espace vide", width="stretch", key="new_space_btn",
                     help="Ouvre un espace neuf. L'ancien reste accessible avec son lien."):
            new_space()


# --------------------------------------------------------------- accueil

def page_home(datasets: list[dict], user: dict) -> None:
    if not datasets:
        masthead("calculBA — espace personnel", "Vos mesures de chantier et de laboratoire, analysées sans tableur.",
                 "Saisissez ou importez les résultats de vos essais : béton, sols, topographie, matériaux, structures, "
                 "routes ou hydraulique. calculBA calcule les statistiques, trace les graphiques et ajuste les lois de "
                 "probabilité. Aucun compte n'est nécessaire et votre espace reste privé.")
        steps = (("01", "Créer un jeu", "Définissez vos colonnes, ou importez un fichier Excel ou CSV."),
                 ("02", "Analyser", "Statistiques, graphiques et lois de probabilité sont calculés automatiquement."),
                 ("03", "Exporter", "Récupérez vos données en Excel ou en CSV, à tout moment."))
        st.markdown('<div class="cb-steps">' + "".join(
            f'<div class="cb-step"><div class="n">{n}</div><b>{t}</b><span>{d}</span></div>' for n, t, d in steps)
            + '</div>', unsafe_allow_html=True)
        b1, b2, _ = st.columns([1.2, 1.2, 2.6])
        if b1.button("Créer un jeu de données", type="primary", width="stretch", key="home_new"):
            go("new")
        if b2.button("Parcourir les exemples", width="stretch", key="home_examples"):
            go("examples")
        st.markdown('<div class="cb-h2">Commencer à partir d\'un exemple</div>'
                    '<p class="cb-sub">Une copie est créée dans votre espace : vous pouvez la modifier ou la supprimer.</p>',
                    unsafe_allow_html=True)
        examples_gallery(user, "home")
        return

    masthead("Espace personnel", "Mes jeux de données")
    total = sum(d["record_count"] for d in datasets)
    stat_strip([("Jeux de données", str(len(datasets))), ("Observations", fmt(total)),
                ("Domaines", str(len({d["domain"] for d in datasets}))),
                ("Dernière modification", datasets[0]["updated_at"][:10])])
    domains = ["Tous"] + sorted({d["domain"] for d in datasets})
    chosen = st.segmented_control("Domaine", domains, default="Tous", key="home_domain", label_visibility="collapsed") or "Tous"
    shown = [d for d in datasets if chosen == "Tous" or d["domain"] == chosen]
    for start in range(0, len(shown), 3):
        cols = st.columns(3)
        for col, d in zip(cols, shown[start:start + 3]):
            with col.container(border=True, key=f"cbcard_ds_{d['id']}"):
                card_title(d["domain"], d["name"])
                st.markdown(f'<div class="cb-desc">{plural(d["record_count"], "observation", "observations")}, '
                            f'{len(d["columns"])} colonnes. Modifié le {d["updated_at"][:10]}.</div>', unsafe_allow_html=True)
                if st.button("Ouvrir", key=f"open_{d['id']}", width="stretch"):
                    go("dataset", ds_id=d["id"], tab="Données", edit_row=None)
    st.caption("Avant toute interprétation, vérifiez les unités de chaque colonne, la qualité des mesures "
               "et les normes d'essai applicables.")


def page_examples(user: dict) -> None:
    masthead("Exemples", "Un jeu de données par domaine",
             "Sept jeux réalistes pour découvrir calculBA dans votre spécialité. « Utiliser » crée une copie dans "
             "votre espace ; « Fichier CSV » montre le format attendu pour importer vos propres mesures.")
    examples_gallery(user, "ex")


# ------------------------------------------------------------ mode d'emploi

def page_guide(user: dict) -> None:
    masthead("Guide", "Mode d'emploi", "calculBA range et analyse vos mesures de génie civil (essais de béton, sols, "
             "topographie, routes, hydraulique…) sans faire les calculs à la main dans un cahier ou un tableur.")
    t1, t2, t3, t4, t5 = st.tabs(["1. Les bases", "2. Mettre des mesures", "3. Lire les résultats",
                                  "4. Mon espace et l'export", "Questions fréquentes"])
    with t1:
        st.subheader("Trois mots à connaître")
        st.markdown(
            "- **Jeu de données** : un ensemble de mesures qui vont ensemble, par exemple « Essais de compression — lot 3 ». "
            "C'est comme une feuille de tableur.\n"
            "- **Colonne** : ce que l'on mesure (Echantillon, Age_jours, Resistance_MPa…).\n"
            "- **Observation** : une ligne, c'est-à-dire une éprouvette, un point ou un sondage, avec toutes ses mesures.")
        st.subheader("Le menu de gauche")
        st.markdown(
            "- **Mes jeux de données** : tous vos jeux, filtrables par domaine.\n"
            "- **Nouveau jeu** : créer un jeu vide ou importer un fichier.\n"
            "- **Exemples par domaine** : un jeu d'exemple pour chaque spécialité (béton, sols, topographie, matériaux, "
            "structures, routes, hydraulique).\n"
            "- **Récents** : vos derniers jeux, en un clic.\n"
            "- **Mon espace** : votre lien personnel, et les boutons pour tout effacer ou repartir de zéro.")
    with t2:
        st.subheader("Méthode A : créer un jeu et saisir les mesures")
        st.markdown(
            "1. Cliquez sur **Nouveau jeu**, puis « Définir les colonnes ».\n"
            "2. Choisissez un **modèle** (Béton, Sols, Topographie…) ou écrivez vos colonnes, **une par ligne**.\n"
            "3. Donnez un nom au jeu, choisissez le domaine, puis **Créer le jeu**.\n"
            "4. Dans l'onglet **Données**, remplissez « Ajouter une observation » et validez. Recommencez pour chaque ligne.")
        st.subheader("Méthode B : importer un fichier Excel ou CSV")
        st.markdown(
            "1. Cliquez sur **Nouveau jeu**, puis « Importer un fichier ».\n"
            "2. **Les noms de colonnes sur la première ligne**, puis une observation par ligne.\n"
            "3. Choisissez le fichier (.xlsx ou .csv), puis **Importer le fichier**.\n"
            "4. Pour voir le bon format, téléchargez le **CSV d'un exemple** dans « Exemples par domaine ».")
        st.subheader("Les bonnes habitudes")
        st.markdown(
            "- **Mettez l'unité dans le nom de la colonne** : `Resistance_MPa`, `Profondeur_m`.\n"
            "- **Virgule ou point** pour les décimales : 30,5 et 30.5 sont acceptés.\n"
            "- **Une mesure inconnue = case vide**, jamais 0 : un 0 ferait baisser la moyenne.\n"
            "- **Une ligne = une éprouvette** (ou un point, un sondage).\n"
            "- Une colonne de **texte** (Lot, Couche, Sondage) sert à **regrouper** les résultats dans les graphiques.")
    with t3:
        st.subheader("Onglet Statistiques")
        st.markdown(
            "| Indicateur | Ce que ça veut dire |\n|---|---|\n"
            "| **n** | nombre de valeurs utilisables dans la colonne |\n"
            "| **Manquantes** | cases vides |\n"
            "| **Moyenne** | la valeur « typique » : somme des valeurs divisée par n |\n"
            "| **Écart-type** | la dispersion : de combien les mesures s'écartent en moyenne de la moyenne |\n"
            "| **CV** | écart-type ÷ moyenne, en % : plus il est petit, plus les mesures sont régulières |\n"
            "| **Min, Max** | plus petite et plus grande valeur : vérifiez qu'il n'y a pas d'erreur de saisie |\n"
            "| **Q1, Médiane, Q3** | 25 %, 50 % et 75 % des valeurs sont en dessous |")
        st.subheader("Onglet Graphiques")
        st.markdown(
            "- **Répartition** : combien de mesures dans chaque intervalle.\n"
            "- **Moyenne par groupe** : compare des lots, des couches, des stations.\n"
            "- **Relation entre deux variables** : la droite et le **R²** disent si deux grandeurs évoluent ensemble.\n"
            "- **Valeurs dans l'ordre de saisie** : repère une dérive ou une valeur anormale.")
        st.subheader("Onglet Probabilités")
        st.markdown(
            "calculBA cherche la **loi de probabilité** qui décrit le mieux vos mesures (Normale, Lognormale, Weibull, "
            "Gumbel…). **La plus petite AIC est la meilleure.** Le **fractile 5 %** est la valeur sous laquelle tombent "
            "5 % des mesures selon la loi. Comptez une trentaine de valeurs pour un résultat fiable.")
        st.caption("Ces calculs aident à interpréter les mesures ; ils ne remplacent pas la vérification "
                   "prévue par les normes (EN 206, Eurocodes…).")
    with t4:
        st.subheader("Un espace privé, sans compte")
        st.markdown(
            "Chaque personne qui ouvre calculBA reçoit **son propre espace vide** : elle ne voit que son travail. "
            "L'espace est lié à **votre lien personnel** (« Mon espace », dans le menu) : gardez-le en favori pour retrouver "
            "vos jeux, sur votre téléphone comme sur votre ordinateur. Ne le partagez qu'avec les personnes qui doivent "
            "voir vos données.")
        st.markdown(
            "- **Tout effacer** : supprime tous vos jeux (Mon espace).\n"
            "- **Commencer un nouvel espace vide** : repart de zéro ; l'ancien reste accessible avec l'ancien lien.\n"
            "- Un espace inutilisé pendant 6 mois est supprimé automatiquement.")
        st.subheader("Exporter")
        st.markdown("Ouvrez le jeu, puis « Exporter, renommer ou supprimer ce jeu » : **Excel (.xlsx)**, "
                    "**CSV Excel FR/DE** (point-virgule, virgule décimale) ou **CSV standard**.")
    with t5:
        st.markdown(
            "**Je ne retrouve plus mes jeux de données.**  \n"
            "Vous avez sans doute ouvert le lien général de calculBA, qui crée un nouvel espace. "
            "Ouvrez votre **lien personnel** (celui qui contient « ?espace= »).\n\n"
            "**L'application affiche « This app has gone to sleep » ?**  \n"
            "Cliquez sur « Yes, get this app back up! » et patientez une minute : les données ne sont pas perdues.\n\n"
            "**« Valeur non numérique » quand j'ajoute une ligne ?**  \n"
            "Écrivez seulement le nombre (30,5) ; l'unité est dans le nom de la colonne.\n\n"
            "**Mon fichier ne s'importe pas.**  \n"
            "Vérifiez que la première ligne contient les noms de colonnes, sans lignes de titre au-dessus, "
            "et que le fichier est en .xlsx ou .csv.")


# ----------------------------------------------------------- nouveau jeu

def page_new(user: dict) -> None:
    st.title("Nouveau jeu de données")
    st.caption("Deux façons de commencer : **définir les colonnes** puis saisir les mesures une par une, "
               "ou **importer un fichier** Excel ou CSV qui contient déjà vos mesures. "
               "Pour voir le bon format, regardez les « Exemples par domaine ».")
    mode = st.segmented_control("Méthode", ["Définir les colonnes", "Importer un fichier"],
                                default="Définir les colonnes", key="new_mode")
    if mode == "Importer un fichier":
        st.write("La première ligne contient les noms de colonnes, chaque ligne suivante une observation. "
                 "Séparateur virgule, point-virgule ou tabulation (détecté automatiquement) ; décimales avec "
                 f"virgule ou point ; {fmt(MAX_ROWS)} lignes et {MAX_COLUMNS} colonnes maximum.")
        st.caption("Un fichier d'exemple pour chaque domaine est disponible dans « Exemples par domaine ».")
        with st.form("import"):
            file = st.file_uploader("Fichier CSV ou Excel (.xlsx)", type=["csv", "txt", "xlsx", "xlsm"])
            name = st.text_input("Nom du jeu (facultatif)", placeholder="Nom du fichier par défaut", max_chars=MAX_NAME_LENGTH)
            domain = st.selectbox("Domaine", DOMAINS, index=DOMAINS.index("Autre"))
            if st.form_submit_button("Importer le fichier", type="primary"):
                if file is None:
                    st.error("Choisissez un fichier CSV ou Excel.")
                else:
                    try:
                        with st.spinner("Import en cours…"):
                            did, n = import_dataset(ensure_space(user), file.name, file.getvalue(), name, domain)
                        flash(f"{plural(n, 'ligne importée', 'lignes importées')}.")
                        go("dataset", ds_id=did, tab="Statistiques", edit_row=None)
                    except AppError as e:
                        st.error(str(e))
        return

    template = st.selectbox("Partir d'un modèle (facultatif)", ["—"] + list(EXAMPLES))
    with st.form("create"):
        name = st.text_input("Nom du jeu", placeholder="Essais de compression — lot 3", max_chars=MAX_NAME_LENGTH)
        domain = st.selectbox("Domaine", DOMAINS, index=DOMAINS.index(template) if template in DOMAINS else 0)
        columns = st.text_area("Colonnes", value=EXAMPLES[template].replace(", ", "\n") if template in EXAMPLES else "",
                               placeholder="Echantillon\nAge_jours\nResistance_MPa",
                               help="Une colonne par ligne ou séparées par des virgules. Indiquez l'unité dans le nom "
                                    "(Resistance_MPa). Le type numérique est détecté automatiquement.")
        if st.form_submit_button("Créer le jeu", type="primary"):
            try:
                did = create_dataset(ensure_space(user), name, domain, columns)
                flash(f"Jeu « {name.strip()} » créé.")
                go("dataset", ds_id=did, tab="Données", edit_row=None)
            except AppError as e:
                st.error(str(e))


# ------------------------------------------------------------- un jeu

TABS = ["Données", "Statistiques", "Graphiques", "Probabilités"]


def page_dataset(user: dict, did: int) -> None:
    with database() as db:
        try:
            ds = get_dataset(db, did, user["id"])
        except AppError as e:
            st.error(str(e))
            if st.button("Revenir à mes jeux de données"):
                go("home")
            return
        records = load_records(db, did)
    rows = [v for _, v in records]
    stats = describe(ds["columns"], rows)

    masthead(ds["domain"], ds["name"], color=domain_color(ds["domain"]))
    stat_strip([("Observations", fmt(ds["record_count"])),
             ("Colonnes (numériques)", f"{len(ds['columns'])} ({len(stats['numeric_columns'])})"),
             ("Modifié le", ds["updated_at"][:10])])

    with st.expander("Exporter, renommer ou supprimer ce jeu"):
        st.caption("Exporter")
        e1, e2, e3 = st.columns(3)
        base = re.sub(r"[^\w\-]+", "_", ds["name"]).strip("_") or f"jeu_{did}"
        e1.download_button("Excel (.xlsx)", lambda: export_bytes(did, user["id"], "xlsx"), file_name=f"calculBA_{base}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
        e2.download_button("CSV Excel FR/DE", lambda: export_bytes(did, user["id"], "csv-excel"), file_name=f"calculBA_{base}.csv",
                           mime="text/csv", width="stretch", help="Point-virgule, décimales avec virgule")
        e3.download_button("CSV standard", lambda: export_bytes(did, user["id"], "csv"), file_name=f"calculBA_{base}_std.csv",
                           mime="text/csv", width="stretch", help="Virgule, décimales avec point")
        st.divider()
        with st.form("rename"):
            c1, c2 = st.columns(2)
            new_name = c1.text_input("Nom", value=ds["name"], max_chars=MAX_NAME_LENGTH)
            new_domain = c2.selectbox("Domaine", DOMAINS, index=DOMAINS.index(ds["domain"]) if ds["domain"] in DOMAINS else 0)
            if st.form_submit_button("Enregistrer le nom et le domaine"):
                try:
                    update_dataset(did, user["id"], new_name, new_domain)
                    flash("Modifications enregistrées.")
                    st.rerun()
                except AppError as e:
                    st.error(str(e))
        st.divider()
        if True:
            sure = st.checkbox(f"Je confirme la suppression définitive de « {ds['name']} » et de ses "
                               f"{plural(ds['record_count'], 'observation', 'observations')}.")
            if st.button("Supprimer le jeu", disabled=not sure):
                try:
                    delete_dataset(did, user["id"])
                    flash(f"Jeu « {ds['name']} » supprimé.")
                    go("home")
                except AppError as e:
                    st.error(str(e))

    tab = st.segmented_control("Section", TABS, key="tab", label_visibility="collapsed") or "Données"
    if tab == "Données":
        tab_data(user, ds, records, stats)
    elif tab == "Statistiques":
        tab_stats(ds, stats)
    elif tab == "Graphiques":
        tab_charts(ds, rows, stats)
    else:
        tab_proba(ds, stats)


def _record_inputs(ds: dict, stats: dict, values: dict, prefix: str) -> dict:
    numeric = set(stats["numeric_columns"])
    cols = st.columns(min(3, max(1, len(ds["columns"]))))
    out = {}
    for i, c in enumerate(ds["columns"]):
        v = values.get(c)
        out[c] = cols[i % len(cols)].text_input(
            c, value="" if v is None else str(v), key=f"{prefix}_{i}", max_chars=MAX_VALUE_LENGTH,
            help="Nombre (virgule ou point)" if c in numeric else None)
    return out


def tab_data(user: dict, ds: dict, records: list, stats: dict) -> None:
    did = ds["id"]
    st.caption("Une ligne = une observation (une éprouvette, un point, un sondage). Laissez vide une mesure inconnue ; "
               "virgule ou point pour les décimales. Les onglets Statistiques, Graphiques et Probabilités se mettent à jour seuls.")
    with st.expander("Ajouter une observation", expanded=not records):
        with st.form("add", clear_on_submit=True):
            values = _record_inputs(ds, stats, {}, "add")
            st.caption("Laissez vide une mesure inconnue plutôt que d'inscrire 0.")
            if st.form_submit_button("Ajouter l'observation", type="primary"):
                try:
                    add_record(did, user["id"], values)
                    flash("Observation ajoutée.")
                    st.rerun()
                except AppError as e:
                    st.error(str(e))

    if not records:
        st.info("Aucune observation. Saisissez une première mesure avec le formulaire ci-dessus.")
        return

    st.subheader(f"Observations ({fmt(len(records))})")
    frame = pd.DataFrame([v for _, v in records], columns=ds["columns"])
    frame.index = range(1, len(frame) + 1)
    frame.index.name = "N°"
    st.dataframe(frame, width="stretch", height=min(420, 38 + 35 * len(frame)))

    st.subheader("Modifier ou supprimer une ligne")
    default_row = st.session_state.get("edit_row") or 1
    row_no = st.number_input("N° de ligne", min_value=1, max_value=len(records), value=min(default_row, len(records)), step=1)
    rid, current = records[row_no - 1]
    with st.form(f"edit_{rid}"):
        st.caption(f"Ligne n° {row_no}")
        new_values = _record_inputs(ds, stats, current, f"edit_{rid}")
        b1, b2 = st.columns([1, 1])
        save = b1.form_submit_button(f"Enregistrer la ligne n° {row_no}", type="primary")
        confirm = b2.checkbox("Confirmer la suppression")
        remove = b2.form_submit_button(f"Supprimer la ligne n° {row_no}")
        if save:
            try:
                update_record(did, user["id"], rid, new_values)
                flash(f"Ligne n° {row_no} modifiée.")
                st.session_state["edit_row"] = row_no
                st.rerun()
            except AppError as e:
                st.error(str(e))
        if remove:
            if not confirm:
                st.error("Cochez « Confirmer la suppression » pour supprimer cette ligne.")
            else:
                try:
                    delete_record(did, user["id"], rid)
                    flash(f"Ligne n° {row_no} supprimée.")
                    st.session_state["edit_row"] = max(1, row_no - 1)
                    st.rerun()
                except AppError as e:
                    st.error(str(e))


def tab_stats(ds: dict, stats: dict) -> None:
    with st.expander("Comment lire ce tableau"):
        st.markdown("**Moyenne** : valeur typique. **Écart-type** : dispersion autour de la moyenne. "
                    "**CV** : écart-type ÷ moyenne, en % (petit = mesures régulières). "
                    "**Min / Max** : vérifiez qu'aucune valeur n'est aberrante. "
                    "**Q1, Médiane, Q3** : 25 %, 50 % et 75 % des valeurs sont en dessous. "
                    "**n** : valeurs utilisées ; **Manquantes** : cases vides (elles ne comptent pas dans les calculs).")
    if not stats["total_rows"]:
        st.info("Ajoutez des observations pour obtenir moyennes, dispersion et répartition.")
        return
    C = stats["columns"]
    flagged = [c for c in stats["numeric_columns"] if C[c]["invalid"]]
    if flagged:
        st.warning("Des valeurs non numériques sont ignorées dans les calculs. Corrigez-les ou videz-les :")
        for c in flagged:
            rows = C[c]["invalid_rows"]
            cols = st.columns([2] + [1] * min(len(rows), 5))
            cols[0].markdown(f"**{c}** : {plural(C[c]['invalid'], 'valeur', 'valeurs')}")
            for j, n in enumerate(rows[:5]):
                if cols[j + 1].button(f"Ligne {n}", key=f"fix_{c}_{n}"):
                    go("dataset", tab="Données", edit_row=n)
    if stats["numeric_columns"]:
        st.subheader("Variables numériques")
        st.dataframe(pd.DataFrame([{
            "Variable": c, "n": fmt(C[c]["count"]), "Manquantes": fmt(C[c]["missing"]),
            "Moyenne": fmt(C[c]["mean"]), "Écart-type": fmt(C[c]["std"]),
            "CV": "—" if C[c]["cv"] is None else fmt(C[c]["cv"], 2) + " %",
            "Min": fmt(C[c]["min"]), "Q1": fmt(C[c]["q1"]), "Médiane": fmt(C[c]["median"]),
            "Q3": fmt(C[c]["q3"]), "Max": fmt(C[c]["max"]),
        } for c in stats["numeric_columns"]]), hide_index=True, width="stretch")
        st.caption(f"Calculées sur {plural(stats['total_rows'], 'observation', 'observations')}. n : valeurs exploitables ; "
                   "écart-type sans biais (n − 1) ; CV : écart-type / |moyenne| ; Q1, Q3 : quartiles.")
    if stats["text_columns"]:
        st.subheader("Variables texte")
        st.dataframe(pd.DataFrame([{
            "Variable": c, "Renseignées": fmt(C[c]["count"]), "Manquantes": fmt(C[c]["missing"]),
            "Distinctes": fmt(C[c]["distinct"]),
            "Valeurs les plus fréquentes": ("Toutes différentes (identifiant)" if C[c]["distinct"] > 1 and C[c]["distinct"] == C[c]["count"]
                                            else ", ".join(f"{t['value']} ({t['count']})" for t in C[c]["top"]) or "—"),
        } for c in stats["text_columns"]]), hide_index=True, width="stretch")


def _chart(chart) -> None:
    st.altair_chart(chart.properties(height=320), width="stretch", theme="streamlit")


def tab_charts(ds: dict, rows: list[dict], stats: dict) -> None:
    nums = stats["numeric_columns"]
    if not stats["total_rows"] or not nums:
        st.info("Rien à tracer pour l'instant : il faut des observations et au moins une colonne numérique.")
        return
    with st.expander("Comment lire ces graphiques"):
        st.markdown("**Répartition** : combien de mesures dans chaque intervalle. "
                    "**Moyenne par groupe** : compare des lots, des âges, des couches. "
                    "**Relation** : chaque point est une observation ; R² proche de 1 = les deux grandeurs évoluent ensemble. "
                    "**Ordre de saisie** : repère une dérive ou une valeur anormale. "
                    "Survolez un graphique pour lire les valeurs.")
    C = stats["columns"]
    by_distinct = sorted(nums, key=lambda c: -C[c]["distinct"])
    y0, x0 = by_distinct[0], by_distinct[-1]

    st.subheader("Répartition")
    col = st.selectbox("Variable", nums, index=nums.index(y0), key="h_col")
    h = C[col]["histogram"]
    # Noms de champs internes simples : un nom de colonne avec apostrophe, point ou
    # crochets (« N° d'essai », « Resistance.MPa ») casserait le graphique Vega-Lite.
    hist = pd.DataFrame({"lo": h["edges"][:-1], "hi": h["edges"][1:], "n": h["counts"]})
    _chart(alt.Chart(hist).mark_bar(color=DATA_COLOR, stroke="white", strokeWidth=1).encode(
        x=alt.X("lo:Q", title=col, bin="binned"), x2="hi:Q", y=alt.Y("n:Q", title="Effectif"),
        tooltip=[alt.Tooltip("lo:Q", title="de"), alt.Tooltip("hi:Q", title="à"), alt.Tooltip("n:Q", title="Effectif")]))
    st.caption(f"{plural(C[col]['count'], 'valeur', 'valeurs')}, moyenne {fmt(C[col]['mean'])}, médiane {fmt(C[col]['median'])}"
               + (f", {plural(C[col]['missing'], 'valeur manquante exclue', 'valeurs manquantes exclues')}" if C[col]["missing"] else ""))

    st.subheader("Moyenne par groupe")
    groupable = [c for c in ds["columns"] if (C[c]["type"] == "text" and 2 <= C[c]["distinct"] <= 60) or (C[c]["type"] == "numeric" and C[c]["count"])]
    default_by = next((c for c in stats["text_columns"] if 2 <= C[c]["distinct"] <= 30), x0)
    g1, g2 = st.columns(2)
    val = g1.selectbox("Variable", nums, index=nums.index(y0), key="g_val")
    by = g2.selectbox("Regrouper par", groupable, index=groupable.index(default_by) if default_by in groupable else 0, key="g_by")
    g = groups(rows, by, val)
    if g["groups"]:
        gf = pd.DataFrame(g["groups"])
        _chart(alt.Chart(gf).mark_bar(color=DATA_COLOR, cornerRadiusTopLeft=4, cornerRadiusTopRight=4).encode(
            x=alt.X("group:N", title=by, sort=list(gf["group"])), y=alt.Y("mean:Q", title=f"Moyenne de {val}"),
            tooltip=[alt.Tooltip("group:N", title=by), alt.Tooltip("mean:Q", title="Moyenne"), alt.Tooltip("n:Q", title="n"),
                     alt.Tooltip("std:Q", title="Écart-type"), alt.Tooltip("min:Q", title="Min"), alt.Tooltip("max:Q", title="Max")]))
        st.caption(f"{plural(g['total_groups'], 'groupe', 'groupes')}" + (f" (les {g['shown']} plus fréquents)" if g["total_groups"] > g["shown"] else ""))

    st.subheader("Relation entre deux variables")
    s1, s2 = st.columns(2)
    x = s1.selectbox("Axe horizontal (X)", nums, index=nums.index(x0), key="s_x")
    y = s2.selectbox("Axe vertical (Y)", nums, index=nums.index(y0), key="s_y")
    sc = scatter(rows, x, y)
    if sc["n"]:
        pts = pd.DataFrame(sc["points"])
        layers = alt.Chart(pts).mark_circle(size=55, color=DATA_COLOR, opacity=0.85).encode(
            x=alt.X("x:Q", title=x, scale=alt.Scale(zero=False)), y=alt.Y("y:Q", title=y, scale=alt.Scale(zero=False)),
            tooltip=[alt.Tooltip("x:Q", title=x), alt.Tooltip("y:Q", title=y)])
        f = sc["fit"]
        if f and x != y:
            line = pd.DataFrame({"x": [f["x_min"], f["x_max"]], "y": [f["slope"] * f["x_min"] + f["intercept"], f["slope"] * f["x_max"] + f["intercept"]]})
            layers = layers + alt.Chart(line).mark_line(color="#888", strokeDash=[6, 4]).encode(x="x:Q", y="y:Q")
        _chart(layers)
        if x == y:
            st.caption("Choisissez deux variables différentes.")
        elif f:
            sign = "−" if f["intercept"] < 0 else "+"
            st.caption(f"Droite des moindres carrés : **{y} = {fmt(f['slope'])} × {x} {sign} {fmt(abs(f['intercept']))}**, "
                       f"R² = {fmt(f['r2'])}, {plural(sc['n'], 'couple', 'couples')}. Une corrélation ne prouve pas une relation de cause à effet.")
        else:
            st.caption("Il faut au moins trois observations où X et Y sont renseignés, et des valeurs qui varient.")

    st.subheader("Valeurs dans l'ordre de saisie")
    scol = st.selectbox("Variable", nums, index=nums.index(y0), key="l_col")
    se = series(rows, scol)
    if se["points"]:
        pts = se["points"]
        filled = []
        for i, p in enumerate(pts):  # coupure de la courbe là où une valeur manque
            if i and p["x"] - pts[i - 1]["x"] > 1 and se["plotted"] == se["available"]:
                filled.append({"x": p["x"] - 0.5, "y": None})
            filled.append(p)
        lf = pd.DataFrame(filled)
        mean = C[scol]["mean"]
        layers = alt.Chart(lf).mark_line(color=DATA_COLOR, point=len(pts) <= 120).encode(
            x=alt.X("x:Q", title="N° d'observation"), y=alt.Y("y:Q", title=scol, scale=alt.Scale(zero=False)),
            tooltip=[alt.Tooltip("x:Q", title="Observation n°"), alt.Tooltip("y:Q", title=scol)])
        if mean is not None:
            layers = layers + alt.Chart(pd.DataFrame({"m": [mean]})).mark_rule(color="#888", strokeDash=[4, 4]).encode(y="m:Q")
        _chart(layers)
        missing = se["total_rows"] - se["available"]
        st.caption(f"{plural(se['available'], 'valeur', 'valeurs')}"
                   + (f", {plural(missing, 'observation sans valeur', 'observations sans valeur')} (coupures)" if missing else "")
                   + ". Trait pointillé : moyenne.")


def tab_proba(ds: dict, stats: dict) -> None:
    C = stats["columns"]
    usable = [c for c in stats["numeric_columns"] if C[c]["count"] >= 5 and C[c]["std"]]
    if not usable:
        st.info("Il faut une colonne numérique avec au moins 5 valeurs différentes pour ajuster une loi de probabilité.")
        return
    usable.sort(key=lambda c: -C[c]["distinct"])
    with st.expander("À quoi sert cet onglet"):
        st.markdown("On cherche la **loi de probabilité** qui décrit le mieux les mesures d'une colonne. "
                    "Barres = vos mesures ; courbe = la loi. Le **fractile 5 %** est la valeur sous laquelle tombent "
                    "5 % des mesures selon cette loi. Dans le tableau du bas, **la plus petite AIC** désigne la loi la plus adaptée. "
                    "Comptez une trentaine de valeurs pour un résultat fiable.")
    st.subheader("Ajuster une loi de probabilité")
    st.caption("Ajustement par maximum de vraisemblance sur toutes les valeurs de la colonne.")
    p1, p2 = st.columns(2)
    col = p1.selectbox("Variable", usable, key="p_col")
    with database() as db:
        values = profile([col], [v for _, v in load_records(db, ds["id"])])[col]["nums"]
    best = fit_distributions(values, "normal")["best"]
    keys = list(DISTRIBUTIONS)
    # Loi proposée : la mieux classée, tant que l'utilisateur n'en choisit pas une autre pour cette variable.
    if st.session_state.get("p_law_for") != col:
        st.session_state["p_law"] = best
        st.session_state["p_law_for"] = col
    law = p2.selectbox("Loi", keys, format_func=lambda k: DISTRIBUTIONS[k]["label"], key="p_law")
    try:
        r = fit_distributions(values, law)
    except FitProblem as e:
        st.error(str(e))
        return
    sel = r["selected"]

    notes = []
    if r["small_sample"]:
        notes.append(f"Seulement {r['n']} valeurs : les paramètres et surtout les fractiles extrêmes sont incertains.")
    notes.append("Le fractile 5 % de la loi ajustée n'est pas directement une valeur caractéristique réglementaire : "
                 "les normes (ex. EN 1990 annexe D, EN 206) ajoutent un facteur lié à la taille de l'échantillon.")
    st.warning(" ".join(notes))

    left, right = st.columns([3, 2])
    with left:
        p = r["plot"]
        bars = pd.DataFrame({"lo": p["edges"][:-1], "hi": p["edges"][1:], "d": p["density"]})
        curve = pd.DataFrame({"x": p["x"], "d": p["pdf"]}).dropna()
        chart = (alt.Chart(bars).mark_bar(color=DATA_COLOR, opacity=0.85, stroke="white", strokeWidth=1).encode(
                     x=alt.X("lo:Q", title=col, bin="binned"), x2="hi:Q", y=alt.Y("d:Q", title="Densité de probabilité"),
                     tooltip=[alt.Tooltip("lo:Q", title="de"), alt.Tooltip("hi:Q", title="à"), alt.Tooltip("d:Q", title="Densité observée")])
                 + alt.Chart(curve).mark_line(color="#222", strokeWidth=2).encode(x="x:Q", y="d:Q"))
        _chart(chart)
        st.caption(f"Barres : mesures ; courbe : loi {sel['label'].lower()} ajustée. {sel['label']} : {sel['use']}")
    with right:
        st.markdown("**Paramètres**")
        st.dataframe(pd.DataFrame([{"Paramètre": q["name"], "Valeur": fmt(q["value"])} for q in sel["parameters"]]
                                  + [{"Paramètre": "Effectif", "Valeur": fmt(r["n"])}]), hide_index=True, width="stretch")
        st.markdown("**Fractiles**")
        label = {0.05: "5 %", 0.5: "50 % (médiane)", 0.95: "95 %"}
        st.dataframe(pd.DataFrame([{"Fractile": label[q["p"]], "Loi ajustée": fmt(q["model"]), "Mesures": fmt(q["empirical"])}
                                   for q in r["quantiles"]]), hide_index=True, width="stretch")
        available = [c for c in r["comparison"] if c["available"]]
        st.caption(f"Test KS indicatif : D = {fmt(sel['ks_d'])}, p = {fmt(sel['ks_p'])}. Rang {sel['rank']} sur {len(available)} selon l'AIC.")

    st.subheader("Quelle loi décrit le mieux ces mesures ?")
    st.dataframe(pd.DataFrame([{
        "Rang": c.get("rank", "—"), "Loi": c["label"] + (" (affichée)" if c["key"] == sel["key"] else ""),
        "AIC": fmt(c["aic"], 2) if c["available"] else "—",
        "Écart": ("meilleure" if c.get("delta_aic") == 0 else "+" + fmt(c["delta_aic"], 2)) if c["available"] else "—",
        "KS : D": fmt(c["ks_d"]) if c["available"] else "—",
        "Usage courant": c["use"] if c["available"] else f"Non applicable : {c['reason']}",
    } for c in r["comparison"]]), hide_index=True, width="stretch")
    st.caption("Classement par critère d'information d'Akaike (AIC) : plus il est petit, meilleur est l'ajustement. "
               "Un écart inférieur à 2 signifie que deux lois décrivent les données presque aussi bien. "
               "Le test de Kolmogorov-Smirnov (KS) est indicatif : les paramètres sont estimés sur ces mêmes données.")


# ------------------------------------------------------------ programme

def main() -> None:
    st.set_page_config(page_title="calculBA — génie civil", layout="wide",
                       initial_sidebar_state="auto")
    inject_css()
    apply_pending()
    show_flash()
    token = current_token()
    user = {"id": find_space(token), "token": token}  # espace privé de ce visiteur
    try:
        purge_idle_spaces()
    except Exception:
        pass  # le ménage ne doit jamais empêcher l'application de s'ouvrir
    datasets = list_datasets(user["id"])
    sidebar(user, datasets)
    view = st.session_state.get("view", "home")
    if view == "guide":
        page_guide(user)
    elif view == "examples":
        page_examples(user)
    elif view == "new":
        page_new(user)
    elif view == "dataset" and st.session_state.get("ds_id"):
        page_dataset(user, st.session_state["ds_id"])
    else:
        page_home(datasets, user)


EXAMPLE_CSV = r'''Echantillon;Lot;Age_jours;Resistance_MPa;Affaissement_mm
B01;Lot A;7;20,6;70
B02;Lot A;7;21,8;70
B03;Lot A;7;21,3;110
B04;Lot A;7;20,4;70
B05;Lot A;7;;70
B06;Lot A;7;20,4;100
B07;Lot A;14;26,0;110
B08;Lot A;14;27,6;100
B09;Lot A;14;28,9;80
B10;Lot A;14;27,7;110
B11;Lot A;14;29,0;100
B12;Lot A;14;26,3;70
B13;Lot A;28;33,0;80
B14;Lot A;28;32,4;90
B15;Lot A;28;30,8;110
B16;Lot A;28;33,5;90
B17;Lot A;28;30,2;70
B18;Lot A;28;31,6;110
B19;Lot B;7;22,6;70
B20;Lot B;7;23,1;110
B21;Lot B;7;23,0;110
B22;Lot B;7;21,5;80
B23;Lot B;7;;90
B24;Lot B;7;23,5;100
B25;Lot B;14;28,0;90
B26;Lot B;14;28,6;80
B27;Lot B;14;30,2;80
B28;Lot B;14;27,1;70
B29;Lot B;14;27,8;90
B30;Lot B;14;28,6;100
B31;Lot B;28;rompu;70
B32;Lot B;28;39,4;110
B33;Lot B;28;32,7;80
B34;Lot B;28;36,3;100
B35;Lot B;28;31,4;70
B36;Lot B;28;36,9;110
B37;Lot C;7;16,6;90
B38;Lot C;7;18,1;90
B39;Lot C;7;18,8;110
B40;Lot C;7;17,5;100
B41;Lot C;7;20,1;90
B42;Lot C;7;19,8;100
B43;Lot C;14;25,3;90
B44;Lot C;14;24,9;110
B45;Lot C;14;28,5;90
B46;Lot C;14;25,4;100
B47;Lot C;14;26,6;100
B48;Lot C;14;24,5;90
B49;Lot C;28;31,4;70
B50;Lot C;28;31,7;80
B51;Lot C;28;31,1;80
B52;Lot C;28;30,2;100
B53;Lot C;28;28,5;70
B54;Lot C;28;33,1;80
'''

# « streamlit run » exécute ce fichier sous le nom __main__ ; un import (tests) ne lance pas l'interface.
if __name__ == "__main__":
    main()
